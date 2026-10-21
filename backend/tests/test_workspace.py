"""Workspace integration against isolated PostgreSQL and synthetic files."""

import io
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import test_api as fixtures
from openpyxl import Workbook, load_workbook


def workbook_bytes():
    book = Workbook()
    cover = book.active
    cover.title = "Титул"
    cover.append(["Synthetic fixture"])
    for name in ("ВОР-1", "ВОР-2"):
        sheet = book.create_sheet(name)
        sheet.append(["Synthetic worksheet"])
        sheet.append(
            [
                "№ п.п.",
                "Наименование работ",
                "Объём работ / Количество",
                "Ед. изм.",
                "Ссылка на ПД",
                "Наименование файла",
                "Номер страниц (через пробел)",
                "Формула",
            ]
        )
        sheet.append([1, "Copper cable", 45, "m", "", "project.pdf", "1", "45"])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class WorkspaceTests(unittest.TestCase):
    _open_app = fixtures.ApiTests._open_app
    _upload = fixtures.ApiTests._upload
    _upload_pair = fixtures.ApiTests._upload_pair

    def test_worker_inactivity_timeout_kills_the_process(self):
        import subprocess
        import sys
        from unittest.mock import Mock
        from uuid import uuid4
        from scripts.audit_worker import stop_worker, wait_for_progress

        repository = Mock()
        repository.run_progress_records.return_value = [{"job_counts": {"running": 1}}]
        with subprocess.Popen(
            [sys.executable, "-c", "import time; time.sleep(30)"],
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
        ) as process:
            try:
                with self.assertRaises(TimeoutError):
                    wait_for_progress(process, repository, uuid4(), 0.1)
                self.assertIsNotNone(process.poll())
                repository.update_run_config.assert_called_once()
            finally:
                stop_worker(process)

    def test_progress_and_history_do_not_load_result_or_job_payloads(self):
        from unittest.mock import patch
        from uuid import UUID

        from workspace_service import run_progress

        with tempfile.TemporaryDirectory() as tmp:
            client, app = self._open_app(tmp)
            vor, pdf = self._upload_pair(client)
            run = self._run(client, app, vor, pdf)
            repo = app.state.context.repository
            record = repo.run_progress_records(run_id=UUID(run))[0]
            project = record["project_id"]
            with (
                patch.object(
                    repo, "load_run", side_effect=AssertionError("heavy read")
                ),
                patch.object(
                    repo, "list_position_jobs", side_effect=AssertionError("heavy read")
                ),
            ):
                progress = client.get(f"/api/v1/audits/{run}/progress")
                self.assertEqual(progress.status_code, 200, progress.text)
                self.assertEqual(progress.json()["done"], 1)
                history = client.get(f"/api/v1/projects/{project}/audits?limit=1")
                self.assertEqual(history.status_code, 200, history.text)
                self.assertEqual(history.json(), [progress.json()])
                self.assertEqual(
                    client.get(
                        f"/api/v1/projects/{project}/audits?offset=1&limit=1"
                    ).json(),
                    [],
                )
            record.update(
                status="running",
                prepare_seconds=None,
                pdf_jobs=1,
                job_counts={"running": 1},
            )
            self.assertEqual(run_progress(app.state.context, record).stage, "preparing")
            record.update(prepare_seconds=1.2, pdf_jobs=0)
            self.assertEqual(run_progress(app.state.context, record).stage, "auditing")
            record["job_counts"] = {"done": 1}
            self.assertEqual(
                run_progress(app.state.context, record).stage, "finalizing"
            )

    def test_production_fallback_runs_in_a_separate_process_without_model_calls(self):
        from dataclasses import replace

        with tempfile.TemporaryDirectory() as tmp:
            client, app = self._open_app(tmp)
            # A missing PDF binding must finish without contacting any provider.
            vor, pdf = self._upload_pair(client)
            repo = app.state.context.repository
            repo_context = replace(app.state.context, client_factory=None)
            app.state.context = repo_context
            # Runner catches child failure using the same repository.
            from api import AuditRunner

            app.state.runner = AuditRunner(repo_context)
            other_pdf = self._upload(
                client,
                "ИГДИ",
                "pdf",
                "unrelated.pdf",
                fixtures._fake_pdf_bytes(),
                "application/pdf",
            )
            run = self._run(client, app, vor, other_pdf)
            progress = client.get(f"/api/v1/audits/{run}/progress")
            self.assertEqual(progress.json()["status"], "done")
            self.assertEqual(
                repo.load_run(__import__("uuid").UUID(run))["results"][0]["requests"], 0
            )

    def test_windows_file_url_writes_into_configured_store(self):
        from pathlib import Path
        from storage_blobs import open_blob_store

        with tempfile.TemporaryDirectory(prefix="blob store ") as tmp:
            store = open_blob_store(f"file://{tmp}")
            ref = store.put(b"configured-store-location")
            self.assertEqual(
                (Path(tmp) / ref.key).read_bytes(), b"configured-store-location"
            )

    def _run(self, client, app, vor, pdf, project="ИГДИ"):
        listed = client.get("/api/v1/projects").json()
        project_id = next(p["id"] for p in listed if p["name"] == project)
        response = client.post(
            "/api/v1/audits",
            json={
                "project_id": project_id,
                "vor_version_id": vor["version_id"],
                "pdf_version_ids": [pdf["version_id"]],
            },
        )
        self.assertEqual(response.status_code, 202, response.text)
        run = response.json()["run_id"]
        self.assertTrue(app.state.runner.wait(__import__("uuid").UUID(run), timeout=60))
        return run

    def test_actual_audit_snapshot_rows_workspace_and_history(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, app = self._open_app(tmp)
            vor, pdf = self._upload_pair(client)
            run = self._run(client, app, vor, pdf)
            progress = client.get(f"/api/v1/audits/{run}/progress")
            self.assertEqual(progress.status_code, 200, progress.text)
            self.assertEqual(progress.json()["status"], "done")
            self.assertEqual(progress.json()["total"], 1)
            summary = client.get(f"/api/v1/audits/{run}/summary")
            self.assertEqual(summary.status_code, 200, summary.text)
            project = summary.json()["project"]["id"]
            rows = client.get(f"/api/v1/audits/{run}/rows").json()
            self.assertEqual(rows["rows"][0]["declared_quantity"], "45")
            self.assertEqual(rows["rows"][0]["queue"], "ok")
            self.assertEqual(rows["rows"][0]["id"], run + ":0")
            view = client.get(f"/api/v1/audits/{run}/positions/0/workspace")
            self.assertEqual(view.status_code, 200, view.text)
            self.assertEqual(
                view.json()["evidence"][0]["version_id"], pdf["version_id"]
            )
            self.assertEqual(view.json()["evidence"][0]["page"], 1)
            history = client.get(f"/api/v1/projects/{project}/audits").json()
            self.assertEqual(history[0]["run_id"], run)
            self.assertEqual(
                client.get(f"/api/v1/audits/{run}/rows?focus=fix").json()["total"], 0
            )
            self.assertEqual(
                client.get(f"/api/v1/audits/{run}/rows?q=Copper").json()["total"], 1
            )
            self.assertEqual(
                client.get(f"/api/v1/audits/{run}/rows?offset=1").json()["rows"], []
            )
            self.assertEqual(
                client.get(f"/api/v1/audits/{run}/positions/99/workspace").status_code,
                404,
            )
            # New input version must not affect a prior run's declared values or sources.
            changed = fixtures._fake_vor_bytes().replace(b";45;", b";99;")
            updated = self._upload(
                client, "ИГДИ", "vor", "vor.csv", changed, "text/csv"
            )
            self.assertNotEqual(updated["version_id"], vor["version_id"])
            self.assertEqual(client.get(f"/api/v1/audits/{run}/rows").json(), rows)
            self.assertEqual(
                client.get(f"/api/v1/audits/{run}/reports/word").status_code, 200
            )
            self.assertEqual(
                client.get(f"/api/v1/audits/{run}/reports/excel").status_code, 400
            )

    def test_content_range_etag_and_native_structure_without_model_calls(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, _ = self._open_app(tmp)
            _, pdf = self._upload_pair(client)
            path = f"/api/v1/documents/{pdf['version_id']}/content"
            original = client.get(path)
            part = client.get(path, headers={"Range": "bytes=0-63"})
            self.assertEqual(part.status_code, 206)
            self.assertEqual(part.content, original.content[:64])
            self.assertEqual(
                client.get(path, headers={"Range": "bytes=-10"}).content,
                original.content[-10:],
            )
            for value in ("bytes=999999-", "bytes=10-1", "bytes=-0", "bytes=0-1,3-4"):
                self.assertEqual(
                    client.get(path, headers={"Range": value}).status_code, 416
                )
            self.assertEqual(
                client.get(
                    path, headers={"If-None-Match": original.headers["etag"]}
                ).status_code,
                304,
            )
            structure = client.get(f"/api/v1/documents/{pdf['version_id']}/structure")
            self.assertEqual(structure.status_code, 200, structure.text)
            self.assertEqual(structure.json()["page_count"], 1)
            self.assertFalse(structure.json()["ocr_dispatched"])

    def test_cross_project_versions_are_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, _ = self._open_app(tmp)
            vor, pdf = self._upload_pair(client)
            other = client.post(
                "/api/v1/projects", json={"name": "Другой проект"}
            ).json()
            body = {
                "project_id": other["id"],
                "vor_version_id": vor["version_id"],
                "pdf_version_ids": [pdf["version_id"]],
            }
            for path in ("documents/preflight", "audits"):
                response = client.post("/api/v1/" + path, json=body)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(
                    response.json()["error"]["code"], "project_version_mismatch"
                )

    def test_every_sheet_keeps_physical_cells_and_preview(self):
        from vor_parse import load_vor_positions

        data = workbook_bytes()
        positions, _ = load_vor_positions(data, "source.xlsx")
        self.assertEqual(
            [p.position_source.sheet for p in positions], ["ВОР-1", "ВОР-2"]
        )
        self.assertEqual([p.index for p in positions], [0, 1])
        self.assertEqual([p.position_source.row for p in positions], [3, 3])
        self.assertIn("C3", positions[0].position_source.cells)
        with tempfile.TemporaryDirectory() as tmp:
            client, app = self._open_app(tmp)
            vor = self._upload(
                client,
                "ИГДИ",
                "vor",
                "source.xlsx",
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            pdf = self._upload(
                client,
                "ИГДИ",
                "pdf",
                "project.pdf",
                fixtures._fake_pdf_bytes(),
                "application/pdf",
            )
            run = self._run(client, app, vor, pdf)
            for index, sheet in enumerate(("ВОР-1", "ВОР-2")):
                view = client.get(f"/api/v1/audits/{run}/positions/{index}/workspace")
                self.assertEqual(view.status_code, 200, view.text)
                self.assertEqual(view.json()["excel"]["sheet"], sheet)
            downloaded = client.get(f"/api/v1/documents/{vor['version_id']}/content")
            book = load_workbook(io.BytesIO(downloaded.content))
            self.assertEqual(book["ВОР-2"]["C3"].value, 45)
            book.close()

    def test_sheet_window_shows_three_rows_and_key_columns(self):
        from vor_parse import load_vor_positions
        from workspace_service import sheet_window

        data = workbook_bytes()
        positions, _ = load_vor_positions(data, "source.xlsx")
        window = sheet_window(data, positions[0])
        rows = [row["number"] for row in window["rows"]]
        # Fixture has a single data row at 3: window clamps to [2, 3].
        self.assertEqual(rows, [2, 3])
        self.assertEqual(window["columns"], ["A", "B", "C", "D", "E", "J", "K"])
        selected = [row for row in window["rows"] if row["selected"]]
        self.assertEqual([row["number"] for row in selected], [3])
        for row in window["rows"]:
            self.assertLessEqual(len(row["cells"]), 7)
    def test_new_routes_require_same_bearer_as_existing_api(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, app = self._open_app(tmp)
            app.state.context.settings.api_token = "synthetic-server-key"
            self.assertEqual(client.get("/api/v1/projects").status_code, 401)
            response = client.get(
                "/api/v1/projects",
                headers={"Authorization": "Bearer synthetic-server-key"},
            )
            self.assertEqual(response.status_code, 200)

    def test_table_quote_highlights_row_with_matching_quantity(self):
        import fitz
        from workspace_service import _quote_boxes

        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72, 100), "2.2 Gruntovye ankery Aker sht. 827")
            page.insert_text((72, 200), "1.2 Gruntovye ankery Aker sht. 231")
            quote = (
                "<tr><td>1.2</td>"
                "<td>Gruntovye ankery Aker</td>"
                "<td>sht.</td><td>231</td></tr>"
            )
            boxes = _quote_boxes(page, quote)
            self.assertEqual(len(boxes), 1)
            self.assertGreater(boxes[0].y0, 150)

    def test_table_quote_without_numbers_stays_unhighlighted(self):
        import fitz
        from workspace_service import _quote_boxes

        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72, 100), "Gruntovye ankery Aker")
            page.insert_text((72, 200), "Gruntovye ankery Aker")
            quote = "<tr><td>Gruntovye ankery Aker</td></tr>"
            self.assertEqual(_quote_boxes(page, quote), [])
