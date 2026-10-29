"""Long-name/long-formula synthetic fixture; only a smeta_test_ui database."""
import io
import json
import os
from pathlib import Path
from uuid import UUID
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from storage_blobs import FileBlobStore
from storage_db import open_repository


def main():
    import sys
    root = Path(sys.argv[1])
    url = os.environ["TEST_POSTGRES_URL"]
    if "/smeta_test_ui" not in url:
        raise SystemExit("Only an isolated E2E database is allowed")
    blobs = FileBlobStore(root / "blobs")
    repo = open_repository(url, blobs)
    old = repo.load_run(UUID((root / "run-id.txt").read_text().strip()))
    original = old["config"]["position_inputs"][0]
    ref = repo.load_version(old["vor_version_id"])
    book = load_workbook(io.BytesIO(blobs.get(ref.blob_key)))
    name = "Устройство нижнего слоя основания из щебеночно-песчаной смеси с уплотнением и уточнением размеров " * 3
    note = "Примечание к выбранной позиции: условия выполнения работ и расчёт расхода материала."
    positions = []
    for i in range(12):
        row = i + 2
        sheet = book.active
        sheet.cell(row, 1, i + 1)
        sheet.cell(row, 2, name)
        sheet.cell(row, 3, 45)
        sheet.cell(row, 4, "m")
        sheet.cell(row, 5, "45*1+" * 20 + "0")
        sheet.cell(row, 9, note)
        sheet.cell(row, 10, "INTERNAL_ID")
        sheet.cell(row, 11, "INTERNAL_TYPE")
        p = json.loads(json.dumps(original))
        p.update(index=i, number=str(i+1), work_name=name, comment=note)
        p["position_source"].update(row=row, cells={f"B{row}": name, f"I{row}": note})
        positions.append(p)
    output = io.BytesIO()
    book.save(output)
    book.close()
    version = repo.add_document_version(old["project_id"], output.getvalue(), filename="layout.xlsx", kind="vor", new_document=True)
    config = {**old["config"], "position_inputs": positions, "prepare_seconds": 0.1}
    run = repo.create_run(old["project_id"], version.id, config,
                          pdf_pins={k: UUID(v) for k, v in config["pdf_versions"].items()})
    result = old["indexed_results"][0]["result"]
    payload = repo.load_position_job(old["id"], 0).payload
    for i in range(12):
        repo.create_position_job(run.id, i, queue="llm", payload=payload)
        repo.save_position_result(run.id, i, result)
        repo.finish_position_job(run.id, i, "done", result=result)
    repo.finish_run(run.id, "done")
    (root / "table-run.json").write_text(json.dumps(dict(run_id=str(run.id))))


if __name__ == "__main__":
    main()
