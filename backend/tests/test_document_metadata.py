"""Tests for document_metadata.extract_document_metadata.

Consumer-visible assertions on the returned dict only: field values,
provenance (sheet/cell/raw), empty-string defaults, and JSON round-trip.
"""

import io
import json
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from document_metadata import extract_document_metadata


def _workbook(cells, merges=(), sheet="ВОР"):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = sheet
    for addr, value in cells.items():
        ws[addr] = value
    for span in merges:
        ws.merge_cells(span)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _table(at_row):
    return {
        f"A{at_row}": "№ п.п.",
        f"B{at_row}": "Наименование работ",
        f"C{at_row}": "Объём работ / Количество",
        f"D{at_row}": "Ед. изм.",
        f"A{at_row + 1}": 1,
        f"B{at_row + 1}": "Укладка плит",
        f"C{at_row + 1}": 10,
        f"D{at_row + 1}": "м3",
    }


def _pdf_bytes(pages):
    """Minimal valid PDF whose native text is exactly ``pages``.

    Non-ASCII characters use byte codes >= 0x80 resolved through an
    embedded ToUnicode CMap, so no font files are needed for Cyrillic.
    ``pages`` is a list of pages; each page is a list of text lines.
    """
    customs = []
    seen = set()
    for page in pages:
        for line in page:
            for char in line:
                if char in seen:
                    continue
                seen.add(char)
                if not 32 <= ord(char) <= 126:
                    customs.append(char)
    if not customs:
        customs.append(" ")
    codes = {char: 128 + index for index, char in enumerate(customs)}

    def encode_line(line):
        out = bytearray()
        for char in line:
            if char in codes:
                out.append(codes[char])
            elif char in "()\\":
                out.extend((ord("\\"), ord(char)))
            else:
                out.append(ord(char))
        return bytes(out)

    page_numbers = [3 + 2 * index for index in range(len(pages))]
    content_numbers = [4 + 2 * index for index in range(len(pages))]
    font_number = 3 + 2 * len(pages)
    cmap_number = 4 + 2 * len(pages)
    last_number = 4 + 2 * len(pages)

    objects = {}
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{number} 0 R" for number in page_numbers)
    objects[2] = (f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>").encode(
        "ascii"
    )
    for index, lines in enumerate(pages):
        if lines:
            parts = []
            height = 780
            for line in lines:
                parts.append(
                    f"BT /F1 12 Tf 50 {height} Td (".encode("ascii")
                    + encode_line(line)
                    + b") Tj ET"
                )
                height -= 14
            stream = b"\n".join(parts)
            objects[content_numbers[index]] = (
                f"<< /Length {len(stream)} >>\nstream\n".encode("ascii")
                + stream
                + b"\nendstream"
            )
            objects[page_numbers[index]] = (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                f"/Resources << /Font << /F1 {font_number} 0 R >> >> "
                f"/Contents {content_numbers[index]} 0 R >>"
            ).encode("ascii")
        else:
            objects[content_numbers[index]] = b"<< /Length 0 >>\nstream\n\nendstream"
            objects[page_numbers[index]] = (
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                b"/Resources << >> >>"
            )
    differences = " ".join(f"{128 + index} /c{index}" for index in range(len(customs)))
    objects[font_number] = (
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
        f"/Encoding << /Differences [{differences}] >> "
        f"/ToUnicode {cmap_number} 0 R >>"
    ).encode("ascii")
    entries = "".join(
        f"<{128 + index:02X}> <{ord(char):04X}>\n" for index, char in enumerate(customs)
    )
    cmap = (
        " /CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n"
        "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) "
        "/Supplement 0 >> def\n"
        "/CMapName /MockRu def\n/CMapType 2 def\n"
        "1 begincodespacerange <00> <FF> endcodespacerange\n"
        f"{len(customs)} beginbfchar\n{entries}endbfchar\n"
        "endcmap\nCMapName currentdict /CMap defineresource pop\nend\nend\n"
    ).encode("ascii")
    objects[cmap_number] = (
        f"<< /Length {len(cmap)} >>\nstream\n".encode("ascii") + cmap + b"endstream"
    )

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = {}
    for number in range(1, last_number + 1):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n".encode("ascii") + objects[number] + b"\nendobj\n"
    startxref = len(out)
    out += f"xref\n0 {last_number + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for number in range(1, last_number + 1):
        out += f"{offsets[number]:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {last_number + 1} /Root 1 0 R >>\n"
        f"startxref\n{startxref}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


class KnownTemplateTests(unittest.TestCase):
    def test_construction_and_work_object_remain_separate(self):
        cells = {
            "B4": "Наименование стройки",
            "C4": "Мост через р. Большая Камышная",
            "B5": "Наименование объекта капитального строительства",
            "C5": "Берегоукрепление откоса на площадке КНС и ЛОС",
            "B6": "Ведомость объемов работ №",
            "C6": "ВР-79/ТКР16",
            "B7": "Основание",
            "C7": "30/09 Раздел №3 Часть №16",
            **_table(12),
        }
        result = extract_document_metadata(_workbook(cells), "renamed.xlsx", "vor")
        self.assertEqual(result["construction_title"], cells["C4"])
        self.assertEqual(result["raw_title"], cells["C5"])
        self.assertEqual(result["document_code"], cells["C6"])
        self.assertEqual(result["pd_reference"], cells["C7"])
        self.assertEqual(
            result["sources"]["raw_title"],
            {"sheet": "ВОР", "cell": "C5", "raw": cells["C5"]},
        )
        self.assertEqual(result["sources"]["document_code"]["cell"], "C6")
        json.dumps(result)

    def test_shifted_merged_values_keep_source_anchor(self):
        cells = {
            "B9": "Наименование стройки",
            "D9": "Мост",
            "B10": "Наименование объекта",
            "D10": "Берегоукрепление",
            "B11": "Номер ведомости",
            "D11": "ВР-79/ТКР16",
            **_table(16),
        }
        data = _workbook(cells, merges=("B9:C9", "D9:F9", "D10:F10"))
        result = extract_document_metadata(data, "vor.xlsx", "vor")
        self.assertEqual(result["construction_title"], "Мост")
        self.assertEqual(result["raw_title"], "Берегоукрепление")
        self.assertEqual(result["document_code"], "ВР-79/ТКР16")
        self.assertEqual(result["sources"]["raw_title"]["cell"], "D10")

    def test_conflicting_declarations_are_not_silently_selected(self):
        cells = {
            "B4": "Основание",
            "C4": "30/09-КР",
            "B5": "Основание",
            "C5": "45/10-АР",
            "B6": "Наименование стройки",
            "C6": "Мост",
            **_table(9),
        }
        result = extract_document_metadata(_workbook(cells), "vor.xlsx", "vor")
        self.assertEqual(result["pd_reference"], "")
        self.assertNotIn("pd_reference", result["sources"])
        self.assertEqual(result["construction_title"], "Мост")

    def test_bare_positional_values_without_labels_ignored(self):
        cells = {
            "C4": "Жилой дом",
            "C5": "Проектный институт",
            "C6": 7,
            "C7": "Служебная записка",
        }
        cells.update(_table(8))
        data = _workbook(cells)
        result = extract_document_metadata(data, "vor.xlsx", "vor")

        self.assertEqual(result["construction_title"], "")
        self.assertEqual(result["raw_title"], "")
        self.assertEqual(result["document_code"], "")
        self.assertEqual(result["pd_reference"], "")
        self.assertEqual(result["sources"], {})

    def test_plain_tabular_workbook_yields_no_metadata(self):
        cells = {
            "A1": "№ п.п.",
            "B1": "Наименование работ",
            "C1": "Объём работ / Количество",
            "D1": "Ед. изм.",
            "E1": "Ссылка на ПД",
            "A2": 1,
            # Must not match the "Основание" label mid-sentence.
            "B2": "Устройство основания из щебня",
            "C2": 100,
            "D2": "м3",
            "E2": "",
        }
        data = _workbook(cells)
        result = extract_document_metadata(data, "vor.xlsx", "vor")

        self.assertEqual(result["construction_title"], "")
        self.assertEqual(result["raw_title"], "")
        self.assertEqual(result["document_code"], "")
        self.assertEqual(result["pd_reference"], "")
        self.assertEqual(result["sources"], {})
        self.assertEqual(result["extractor_version"], "1")


class CsvMetadataTests(unittest.TestCase):
    def test_labelled_csv_preamble_is_extracted_before_work_table(self):
        for encoding in ("utf-8-sig", "cp1251"):
            payload = (
                "Наименование стройки;Мост\n"
                "Наименование объекта;Берегоукрепление\n"
                "Номер ведомости;ВР-79/ТКР16\n"
                "Основание;30/09-КР\n"
                "№ п.п.;Наименование работ;Количество;Ед. изм.\n"
                "1;Основание;45;м3\n"
            ).encode(encoding)
            result = extract_document_metadata(payload, "vor.csv", "vor")
            self.assertEqual(result["construction_title"], "Мост")
            self.assertEqual(result["raw_title"], "Берегоукрепление")
            self.assertEqual(result["document_code"], "ВР-79/ТКР16")
            self.assertEqual(result["pd_reference"], "30/09-КР")
            self.assertEqual(result["sources"]["raw_title"]["cell"], "B2")

    def test_csv_with_label_like_headers_yields_no_metadata(self):
        payload = (
            "№ п.п.;Наименование работ;Объём работ / Количество;Ед. изм.;"
            "Ссылка на ПД;Наименование файла;Номер страниц (через пробел)\n"
            "1;Медный кабель;45;м;;project.pdf;1\n"
        ).encode("utf-8-sig")
        result = extract_document_metadata(payload, "vor.csv", "vor")

        self.assertEqual(result["construction_title"], "")
        self.assertEqual(result["raw_title"], "")
        self.assertEqual(result["document_code"], "")
        self.assertEqual(result["pd_reference"], "")
        self.assertEqual(result["sources"], {})

    def test_csv_without_headers_yields_no_metadata(self):
        payload = "1;Медный кабель;45\n2;Укладка плит;10\n".encode("cp1251")
        result = extract_document_metadata(payload, "vor.csv", "vor")

        self.assertEqual(result["construction_title"], "")
        self.assertEqual(result["raw_title"], "")
        self.assertEqual(result["document_code"], "")
        self.assertEqual(result["pd_reference"], "")
        self.assertEqual(result["sources"], {})
        self.assertEqual(result["extractor_version"], "1")


class PdfMetadataTests(unittest.TestCase):
    def test_native_title_block_metadata(self):
        data = _pdf_bytes(
            [
                ["Шифр: 123-КР", "Объект: Жилой дом"],
                ["Наименование документа: Конструкции железобетонные"],
                [],
            ]
        )
        result = extract_document_metadata(data, "pd.pdf", "pdf")

        self.assertEqual(result["document_code"], "123-КР")
        self.assertEqual(result["construction_title"], "Жилой дом")
        self.assertEqual(result["raw_title"], "Конструкции железобетонные")
        self.assertEqual(result["pd_reference"], "")
        self.assertEqual(
            result["sources"]["document_code"],
            {"sheet": "page 1", "cell": "L1", "raw": "Шифр: 123-КР"},
        )
        self.assertEqual(
            result["sources"]["construction_title"],
            {"sheet": "page 1", "cell": "L2", "raw": "Объект: Жилой дом"},
        )
        self.assertEqual(
            result["sources"]["raw_title"],
            {
                "sheet": "page 2",
                "cell": "L1",
                "raw": "Наименование документа: Конструкции железобетонные",
            },
        )
        self.assertEqual(result["extractor_version"], "1")
        json.dumps(result)

    def test_metadata_bounded_to_first_three_pages(self):
        data = _pdf_bytes(
            [
                ["Служебная записка"],
                [],
                [],
                ["Наименование документа: Поздний заголовок"],
            ]
        )
        result = extract_document_metadata(data, "pd.pdf", "pdf")

        # Designated title beyond page 3 must not be picked up.
        self.assertEqual(result["raw_title"], "")
        self.assertEqual(result["document_code"], "")
        self.assertEqual(result["construction_title"], "")
        self.assertEqual(result["pd_reference"], "")
        self.assertEqual(result["sources"], {})

    def test_blank_pdf_yields_no_metadata(self):
        result = extract_document_metadata(_pdf_bytes([[]]), "pd.pdf", "pdf")

        self.assertEqual(result["construction_title"], "")
        self.assertEqual(result["raw_title"], "")
        self.assertEqual(result["document_code"], "")
        self.assertEqual(result["pd_reference"], "")
        self.assertEqual(result["sources"], {})

    def test_unreadable_payloads_yield_no_metadata(self):
        for payload, filename, kind in (
            (b"definitely not a pdf", "pd.pdf", "pdf"),
            (b"not a workbook", "vor.xlsx", "vor"),
        ):
            result = extract_document_metadata(payload, filename, kind)
            self.assertEqual(result["construction_title"], "")
            self.assertEqual(result["raw_title"], "")
            self.assertEqual(result["document_code"], "")
            self.assertEqual(result["pd_reference"], "")
            self.assertEqual(result["sources"], {})
            self.assertEqual(result["extractor_version"], "1")


if __name__ == "__main__":
    unittest.main()
