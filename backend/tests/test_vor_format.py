"""VOR formatting checks: first sheet, filtered issues, persisted verdict."""

import io
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import test_api as fixtures
from openpyxl import Workbook
from openpyxl.styles import Font

HEADERS = [
    "№ п.п.",
    "Наименование работ, ресурсов, затрат по проекту",
    "Ед. изм.",
    "Объем работ / Количество",
    "Формула расчета объемов работ и расхода материалов, потребности ресурсов",
    "Ссылка на чертежи, спецификации в проектной документации",
    "Наименование файла",
    "Номера страниц (через пробел)",
    "Дополнительная информация (комментарий).",
    "Идентификатор",
    "Тип позиции",
]


def _template_book(second_sheet=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "ВР-79 ТКР16"
    sheet["A1"] = "Документ"
    sheet["D1"] = "Ведомость объемов работ"
    sheet["A2"] = "Версия"
    sheet["D2"] = "3_01"
    sheet["A4"] = "Наименование стройки"
    sheet["D4"] = "Мост через реку"
    sheet["A5"] = "Наименование объекта капитального строительства"
    sheet["D5"] = "Берегоукрепление откоса"
    sheet["A6"] = "Ведомость объемов работ №"
    sheet["D6"] = "В.79.16"
    sheet["A7"] = "Основание((наименование раздела (подраздела) ПД))"
    sheet["D7"] = "Раздел ПД № 3"
    sheet["A8"] = "Дата составления"
    sheet["D8"] = "24.03.2025"
    sheet["A10"] = "Составил ФИО"
    sheet["D10"] = "Иванова И.И."
    sheet["A11"] = "Составил должность"
    sheet["D11"] = "Инженер"
    sheet["A12"] = "Проверил ФИО"
    sheet["D12"] = "Петров П.П."
    sheet["A13"] = "Проверил должность"
    sheet["D13"] = "ГИП"
    for col, header in enumerate(HEADERS, start=1):
        sheet.cell(row=15, column=col, value=header)
    for col, number in enumerate(
        ["1", "2", "3", "4", "5", "6", "6.1", "6.2", "7", "8", "9"], start=1
    ):
        sheet.cell(row=16, column=col, value=number)
    sheet["A17"] = "Раздел: 1. Земляные работы"
    sheet["A17"].font = Font(bold=True)
    sheet.append(
        [
            1,
            "Берегоукрепление",
            "м3",
            100,
            "",
            "30/09-ТКР16.СО л. 1",
            "design.pdf",
            "1",
            "",
            "Ф78",
            "Работа",
        ]
    )
    if second_sheet:
        other = book.create_sheet("Приложение")
        other["A1"] = "мусор"
        book.active = book.index(other)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class VorFormatTests(unittest.TestCase):
    def test_first_sheet_used_without_name(self):
        from vor_rules import validate_vor_formatting

        rows = validate_vor_formatting(io.BytesIO(_template_book(second_sheet=True)))
        bad = [row for row in rows if row["status"] != "ОК"]
        self.assertEqual(bad, [])

    def test_blank_sheet_name_means_first_sheet(self):
        from vor_rules import validate_vor_formatting

        for name in (None, "", "   "):
            rows = validate_vor_formatting(
                io.BytesIO(_template_book(second_sheet=True)), sheet_name=name
            )
            bad = [row for row in rows if row["status"] != "ОК"]
            self.assertEqual(bad, [], name)

    def test_unknown_sheet_raises_key_error(self):
        from vor_rules import validate_vor_formatting

        with self.assertRaises(KeyError):
            validate_vor_formatting(
                io.BytesIO(_template_book()), sheet_name="Нет такого"
            )

    def test_broken_cells_report_human_readable_rows(self):
        from vor_rules import validate_vor_formatting

        book = Workbook()
        sheet = book.active
        sheet.title = "ВР"
        sheet["A1"] = "Документ"
        output = io.BytesIO()
        book.save(output)
        book.close()
        output.seek(0)
        rows = validate_vor_formatting(output)
        bad = [row for row in rows if row["status"] != "ОК"]
        self.assertTrue(bad)
        for row in bad:
            self.assertTrue(row["coord"])
            self.assertTrue(row["error"])
            self.assertNotIn("False", row["error"])
            self.assertNotIn("None", row["error"])
        coords = {row["coord"] for row in bad}
        self.assertIn("D1", coords)

    def test_summary_orders_positions_first_and_groups_stray_columns(self):
        from vor_rules import summarize_format_rows

        rows = [
            {
                "coord": "L18",
                "type": "COL_L_LIMIT",
                "status": "Ошибка",
                "expected": "x",
                "actual": "1",
                "error": "y",
            },
            {
                "coord": "M18",
                "type": "COL_L_LIMIT",
                "status": "Ошибка",
                "expected": "x",
                "actual": "2",
                "error": "y",
            },
            {
                "coord": "D1",
                "type": "STATIC",
                "status": "Ошибка",
                "expected": "x",
                "actual": "",
                "error": "y",
            },
            {
                "coord": "E20",
                "type": "ROW_POS",
                "status": "Ошибка",
                "expected": "x",
                "actual": "1",
                "error": "y",
            },
            {
                "coord": "A17",
                "type": "SECTION",
                "status": "Ошибка",
                "expected": "x",
                "actual": "1",
                "error": "y",
            },
            {
                "coord": "A1",
                "type": "STATIC",
                "status": "ОК",
                "expected": "x",
                "actual": "1",
                "error": "",
            },
        ]
        issues, errors = summarize_format_rows(rows)
        self.assertEqual(errors, 5)
        self.assertEqual(
            [row["coord"] for row in issues],
            ["E20", "A17", "D1", "L18", "M18"],
        )

    def test_stray_columns_collapse_to_one_row(self):
        from vor_rules import validate_vor_formatting

        book = Workbook()
        sheet = book.active
        sheet.title = "ВР"
        sheet["A1"] = "Документ"
        for col, header in enumerate(HEADERS, start=1):
            sheet.cell(row=15, column=col, value=header)
        sheet["L18"] = 526
        sheet["M22"] = 1260
        output = io.BytesIO()
        book.save(output)
        book.close()
        output.seek(0)
        rows = validate_vor_formatting(output)
        stray = [row for row in rows if row["type"] == "COL_L_LIMIT"]
        self.assertEqual(len(stray), 1)
        self.assertIn("L18", stray[0]["coord"])
        self.assertIn("M22", stray[0]["coord"])
        self.assertIn("2 ячеек", stray[0]["actual"])

    def test_user_strings_have_no_regex(self):
        from vor_rules import validate_vor_formatting

        book = Workbook()
        sheet = book.active
        sheet.title = "ВР"
        sheet["A1"] = "Документ"
        output = io.BytesIO()
        book.save(output)
        book.close()
        output.seek(0)
        for row in validate_vor_formatting(output):
            for key in ("expected", "error"):
                self.assertNotIn("(^", str(row.get(key) or ""))
                self.assertNotIn("\\d", str(row.get(key) or ""))

    def test_summary_assigns_group_titles(self):
        from vor_rules import summarize_format_rows

        rows = [
            {
                "coord": "J20",
                "type": "ROW_POS",
                "status": "Ошибка",
                "expected": "x",
                "actual": "",
                "error": "Идентификатор не заполнен",
            },
            {
                "coord": "J21",
                "type": "ROW_POS",
                "status": "Ошибка",
                "expected": "x",
                "actual": "",
                "error": "Идентификатор не заполнен",
            },
            {
                "coord": "D1",
                "type": "STATIC",
                "status": "Ошибка",
                "expected": "x",
                "actual": "",
                "error": "Должно быть ровно 'Документ'",
            },
        ]
        issues, _ = summarize_format_rows(rows)
        by_coord = {row["coord"]: row for row in issues}
        self.assertEqual(by_coord["J20"]["group"], by_coord["J21"]["group"])
        self.assertEqual(by_coord["J20"]["group_title"], "Не заполнен идентификатор")
        self.assertTrue(by_coord["J20"]["group_hint"])
        self.assertEqual(by_coord["D1"]["group"], "HEADER")

    def test_highlight_paints_cells_with_comments(self):
        import openpyxl
        from vor_rules import get_highlighted_workbook

        painted = get_highlighted_workbook(
            io.BytesIO(_template_book()),
            [
                {
                    "coord": "D1",
                    "type": "STATIC",
                    "status": "Ошибка",
                    "expected": "Впишите X",
                    "actual": "",
                    "error": "Не заполнено",
                },
                {
                    "coord": "L18, M18",
                    "type": "COL_L_LIMIT",
                    "status": "Ошибка",
                    "expected": "Уберите",
                    "actual": "2",
                    "error": "Мусор справа",
                },
            ],
            sheet_name="ВР-79 ТКР16",
        )
        wb = openpyxl.load_workbook(io.BytesIO(painted))
        ws = wb["ВР-79 ТКР16"]
        self.assertEqual(ws["D1"].fill.start_color.rgb, "00FCE8E9")
        self.assertIn("Не заполнено", ws["D1"].comment.text)
        self.assertIn("Как исправить", ws["D1"].comment.text)
        self.assertEqual(ws["L18"].fill.start_color.rgb, "00FCE8E9")
        self.assertEqual(ws["M18"].fill.start_color.rgb, "00FCE8E9")

    def test_word_report_has_same_columns_and_groups(self):
        import docx
        from vor_rules import generate_rules_word_report

        report = generate_rules_word_report(
            "vor.xlsx",
            [
                {
                    "coord": "J20",
                    "type": "ROW_POS",
                    "status": "Ошибка",
                    "expected": "Впишите Ф78",
                    "actual": "",
                    "error": "Идентификатор не заполнен",
                },
            ],
            sheet="ВР-79 ТКР16",
        )
        doc = docx.Document(io.BytesIO(report))
        texts = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("vor.xlsx", texts)
        self.assertIn("ВР-79 ТКР16", texts)
        self.assertIn("Не заполнен идентификатор", texts)
        self.assertEqual(len(doc.tables), 1)
        headers = [cell.text for cell in doc.tables[0].rows[0].cells]
        self.assertEqual(headers, ["Ячейка", "Что не так", "Как исправить"])
        first = [cell.text for cell in doc.tables[0].rows[1].cells]
        self.assertEqual(first[0], "J20")


class FormatApiTests(unittest.TestCase):
    _open_app = fixtures.ApiTests._open_app
    _upload = fixtures.ApiTests._upload

    def _project_id(self, client, name="Формат"):
        response = client.post("/api/v1/projects", json={"name": name})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()["id"]

    def test_check_filters_ok_and_persists_passed(self):
        with tempfile.TemporaryDirectory(prefix="smeta-format-") as tmp:
            client, _ = self._open_app(tmp)
            project_id = self._project_id(client)
            uploaded = self._upload(
                client,
                "Формат",
                "vor",
                "vor.xlsx",
                _template_book(second_sheet=True),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                project_id=project_id,
                new_document=True,
            )
            response = client.post(
                f"/api/v1/documents/{uploaded['version_id']}/format-checks"
            )
            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertEqual(payload["sheet"], "ВР-79 ТКР16")
            self.assertEqual(payload["sheets"], ["ВР-79 ТКР16", "Приложение"])
            self.assertEqual(payload["issues"], [])
            self.assertEqual(payload["error_count"], 0)
            self.assertEqual(payload["warning_count"], 0)
            detail = client.get(
                f"/api/v1/documents/{uploaded['document_id']}/detail"
            ).json()
            status = detail["versions"][0]["format_status"]
            self.assertEqual(status["state"], "passed")
            self.assertEqual(status["sheet"], "ВР-79 ТКР16")
            self.assertIn("checked_at", status)

    def test_check_failed_persists_counts(self):
        with tempfile.TemporaryDirectory(prefix="smeta-format-") as tmp:
            client, _ = self._open_app(tmp)
            project_id = self._project_id(client)
            book = Workbook()
            sheet = book.active
            sheet.title = "ВР"
            sheet["A1"] = "Документ"
            for col, header in enumerate(HEADERS, start=1):
                sheet.cell(row=15, column=col, value=header)
            sheet.append([1, "Работа", "м3", 10, "", "", "", "1", "", "Ф1", "Работа"])
            output = io.BytesIO()
            book.save(output)
            book.close()
            uploaded = self._upload(
                client,
                "Формат",
                "vor",
                "broken.xlsx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                project_id=project_id,
                new_document=True,
            )
            payload = client.post(
                f"/api/v1/documents/{uploaded['version_id']}/format-checks"
            ).json()
            self.assertGreater(payload["error_count"], 0)
            self.assertTrue(payload["issues"])
            self.assertNotIn("ОК", {row["status"] for row in payload["issues"]})
            detail = client.get(
                f"/api/v1/documents/{uploaded['document_id']}/detail"
            ).json()
            status = detail["versions"][0]["format_status"]
            self.assertEqual(status["state"], "failed")
            self.assertEqual(status["error_count"], payload["error_count"])

    def test_unknown_sheet_is_404(self):
        with tempfile.TemporaryDirectory(prefix="smeta-format-") as tmp:
            client, _ = self._open_app(tmp)
            project_id = self._project_id(client)
            uploaded = self._upload(
                client,
                "Формат",
                "vor",
                "vor.xlsx",
                _template_book(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                project_id=project_id,
                new_document=True,
            )
            response = client.post(
                f"/api/v1/documents/{uploaded['version_id']}/format-checks"
                "?sheet=Нет%20такого"
            )
            self.assertEqual(response.status_code, 404)
            self.assertEqual(response.json()["error"]["code"], "sheet_not_found")

    def test_standalone_check_does_not_store_documents(self):
        with tempfile.TemporaryDirectory(prefix="smeta-format-") as tmp:
            client, _ = self._open_app(tmp)
            project_id = self._project_id(client)
            before = client.get("/api/v1/projects").json()
            response = client.post(
                "/api/v1/format-checks",
                files={
                    "file": (
                        "probe.xlsx",
                        _template_book(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertEqual(payload["version_id"], "")
            self.assertEqual(payload["sheet"], "ВР-79 ТКР16")
            self.assertEqual(payload["issues"], [])
            after = client.get("/api/v1/projects").json()
            self.assertEqual(before, after)
            catalog = client.get(f"/api/v1/projects/{project_id}/catalog").json()
            self.assertEqual(catalog, [])

    def test_standalone_check_rejects_non_xlsx(self):
        with tempfile.TemporaryDirectory(prefix="smeta-format-") as tmp:
            client, _ = self._open_app(tmp)
            response = client.post(
                "/api/v1/format-checks",
                files={"file": ("vor.csv", b"a;b\n", "text/csv")},
            )
            self.assertEqual(response.status_code, 400)
            self.assertEqual(response.json()["error"]["code"], "xlsx_required")

    def test_format_reports_download(self):
        with tempfile.TemporaryDirectory(prefix="smeta-format-") as tmp:
            client, _ = self._open_app(tmp)
            project_id = self._project_id(client)
            uploaded = self._upload(
                client,
                "Формат",
                "vor",
                "vor.xlsx",
                _template_book(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                project_id=project_id,
                new_document=True,
            )
            version_id = uploaded["version_id"]
            excel = client.get(f"/api/v1/documents/{version_id}/format-report")
            self.assertEqual(excel.status_code, 200, excel.text[:200])
            self.assertTrue(excel.content.startswith(b"PK"))
            word = client.get(f"/api/v1/documents/{version_id}/format-word")
            self.assertEqual(word.status_code, 200, word.text[:200])
            self.assertTrue(word.content.startswith(b"PK"))
            import docx

            doc = docx.Document(io.BytesIO(word.content))
            self.assertIn(
                "соответствует шаблону",
                "\n".join(p.text for p in doc.paragraphs),
            )

    def test_standalone_reports_download_without_storing(self):
        with tempfile.TemporaryDirectory(prefix="smeta-format-") as tmp:
            client, _ = self._open_app(tmp)
            project_id = self._project_id(client)
            files = {
                "file": (
                    "probe.xlsx",
                    _template_book(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            }
            excel = client.post("/api/v1/format-report", files=files)
            self.assertEqual(excel.status_code, 200, excel.text[:200])
            self.assertTrue(excel.content.startswith(b"PK"))
            word = client.post("/api/v1/format-word", files=files)
            self.assertEqual(word.status_code, 200, word.text[:200])
            self.assertTrue(word.content.startswith(b"PK"))
            catalog = client.get(f"/api/v1/projects/{project_id}/catalog").json()
            self.assertEqual(catalog, [])


if __name__ == "__main__":
    unittest.main()
