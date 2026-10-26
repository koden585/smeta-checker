"""User regressions: paired schedule units, nonsemantic A/J, optional formula."""

import io
import json
from decimal import Decimal

import openpyxl
import pytest

from audit_core import _bind_quantity_fact, build_claim_specs, parse_audit_response
from audit_measurements import (
    iter_paired_measurements,
    paired_measurements,
    paired_units,
)
from audit_service import _independent_checks
from vor_formula import has_formula_text
from vor_parse import load_vor_positions
from vor_rules import validate_vor_formatting
from test_vor_format import _template_book


def audit(quote, volume="114", unit="шт.", work="Шпала ШТУ 03.02 со скреплениями"):
    return parse_audit_response(
        json.dumps(
            {
                "claims": [
                    {"claim_id": "presence", "source_ids": ["row"]},
                    {"claim_id": "quantity", "source_ids": ["row"]},
                    {"claim_id": "unit", "source_ids": ["row"]},
                ]
            }
        ),
        {24: quote},
        False,
        expected_volume=volume,
        expected_unit=unit,
        expected_work=work,
        position_type="Материал",
        quote_sources={"row": {"page": 24, "quote": quote}},
    )


@pytest.mark.parametrize(
    "row",
    [
        "10.2 | Шпала ШТУ 03.02 со скреплениями | шт./т | 114/23,53",
        "10.2 Шпала ШТУ 03.02 со скреплениями (3) шт./т 114/23,53",
        "Шпала ШТУ 03.02 со скреплениями: 114/23.53 шт./т",
        "<table><tr><td>10.2</td><td>Шпала ШТУ 03.02 со скреплениями</td><td>шт./т</td><td>114/23,53</td></tr></table>",
        "Шпала ШТУ 03.02 со скреплениями: 114 шт./23,53 т",
    ],
)
def test_sleeper_screenshot_quantity_and_unit_confirmed(row):
    result = audit(row)
    assert result["volume_confirmed"] is True
    assert result["unit_confirmed"] is True
    assert result["found_volume"] == "114"
    assert result["found_unit"] == "шт"


@pytest.mark.parametrize(
    "unit,expected,actual,actual_unit",
    [
        ("шт.", "114", "114", "шт"),
        ("т", "23,53", "23.53", "т"),
        ("кг", "23530", "23.53", "т"),
    ],
)
def test_each_side_selected_by_unit_not_value(unit, expected, actual, actual_unit):
    result = audit("Шпала ШТУ 03.02 | шт./т | 114/23,53", expected, unit)
    assert result["volume_confirmed"]
    assert result["found_volume"] == actual and result["found_unit"] == actual_unit


def test_reversed_order_uses_second_count():
    result = audit("Шпала ШТУ 03.02 | т/шт. | 23,53/114")
    assert result["volume_confirmed"] and result["found_volume"] == "114"


def test_wrong_quantity_is_not_hidden_by_a_matching_other_number():
    result = audit("Шпала ШТУ 03.02 | шт./т | 115/114")
    check = next(c for c in result["checks"] if c["claim_id"] == "quantity")
    assert check["state"] == "contradiction"
    assert result["found_volume"] == "115" and not result["volume_confirmed"]


def test_embedded_rail_length_does_not_replace_table_amount():
    result = audit(
        "Рельсы РТ62 длиной 12,5 м | м/т | 124,62/7,74", "124,62", "м", "Рельсы РТ62"
    )
    assert result["volume_confirmed"] and result["found_volume"] == "124.62"


@pytest.mark.parametrize(
    "text",
    [
        "Шпала 114/23,53",
        "Шпала | шт./т | 114",
        "Шпала | шт./т | 114/23,53/5",
        "Шпала | шт./шт. | 114/115",
    ],
)
def test_missing_or_ambiguous_unit_mapping_is_not_guessed(text):
    fact, _ = _bind_quantity_fact([text], expected_work="Шпала", expected_unit="шт.")
    assert fact is None


@pytest.mark.parametrize("unit", ["кг/м2", "г/м2", "м3/м2", "кН/м", "л/с"])
def test_dimensional_rates_are_not_split_into_alternative_units(unit):
    assert paired_units(unit) == ()
    assert paired_measurements("114/23,53", unit) == []
    assert iter_paired_measurements("0,5 " + unit) == []


@pytest.mark.parametrize(
    "units,values",
    [("м2/т", "631,2/0,165"), ("м3/т", "20/35"), ("шт./т", "1 114/23,53")],
)
def test_supported_area_volume_mass_and_thousands(units, values):
    measurements = iter_paired_measurements(units + " " + values)
    assert len(measurements) == 2
    assert all(m.value is not None and m.dimension for m in measurements)


def test_explicit_cells_join_only_with_same_row_provenance():
    context = {24: "Шпала ШТУ 03.02 | шт./т | 114/23,53"}
    sources = {
        "work": {"page": 24, "quote": "Шпала ШТУ 03.02"},
        "value": {
            "page": 24,
            "quote": "114/23,53",
            "locator": {"table_id": "t1", "row_id": "r2"},
        },
        "units": {
            "page": 24,
            "quote": "шт./т",
            "locator": {"table_id": "t1", "row_id": "r2"},
        },
    }
    raw = json.dumps(
        {
            "claims": [
                {"claim_id": "presence", "source_ids": ["work"]},
                {
                    "claim_id": "quantity",
                    "source_ids": ["value", "units"],
                    "value_source_id": "value",
                    "unit_source_id": "units",
                },
                {"claim_id": "unit", "source_ids": ["units"]},
            ]
        }
    )

    def evaluate():
        return parse_audit_response(
            raw,
            context,
            False,
            expected_volume="114",
            expected_unit="шт.",
            expected_work="Шпала",
            position_type="Материал",
            quote_sources=sources,
        )

    assert evaluate()["volume_confirmed"]
    sources["units"]["locator"]["row_id"] = "other-row"
    assert not evaluate()["volume_confirmed"]


def workbook(**changes):
    book = openpyxl.load_workbook(io.BytesIO(_template_book()))
    sheet = book.active
    for address, value in changes.items():
        sheet[address] = value
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


@pytest.mark.parametrize(
    "number", [None, "", "любая метка", "Раздел произвольный", "Заголовок"]
)
def test_service_number_and_missing_identifier_do_not_reject_a_work(number):
    data = workbook(A18=number, J18=None)
    rows = validate_vor_formatting(io.BytesIO(data))
    assert not [r for r in rows if r["status"] != "ОК"]
    positions, _ = load_vor_positions(data, "vor.xlsx")
    assert len(positions) == 1 and positions[0].work_name == "Берегоукрепление"
    assert not any(
        c.startswith(("A", "J"))
        for spec in build_claim_specs(positions[0])
        for c in spec.cells
    )


def test_duplicate_service_identifiers_are_not_errors():
    book = openpyxl.load_workbook(io.BytesIO(_template_book()))
    sheet = book.active
    for col in range(1, 12):
        sheet.cell(19, col, sheet.cell(18, col).value)
    stream = io.BytesIO()
    book.save(stream)
    rows = validate_vor_formatting(io.BytesIO(stream.getvalue()))
    assert not [r for r in rows if r["status"] != "ОК"]


@pytest.mark.parametrize(
    "formula", [None, "", " ", "-", "—", "не задана", "не задано", "не требуется"]
)
def test_omitted_formula_is_not_an_error(formula):
    data = workbook(E18=formula)
    assert not has_formula_text(formula)
    assert not [
        r for r in validate_vor_formatting(io.BytesIO(data)) if r["status"] != "ОК"
    ]
    positions, _ = load_vor_positions(data, "vor.xlsx")
    assert positions[0].formula == "" and Decimal(positions[0].volume) == Decimal("100")
    assert not any(
        c["claim_id"] == "workbook:formula"
        for c in _independent_checks({"calculation": {}}, [])
    )


def test_existing_bad_formula_still_checked():
    assert has_formula_text(0) and has_formula_text("50+20")
    rows = validate_vor_formatting(io.BytesIO(workbook(E18="50+20")))
    assert any(r["status"] == "Ошибка" and "формул" in str(r).lower() for r in rows)
    checks = _independent_checks(
        {"calculation": {"status": "error", "expression": "=1/0"}}, []
    )
    assert checks[0]["state"] == "contradiction"


def test_optional_calculation_placeholder_does_not_create_required_check():
    assert (
        _independent_checks(
            {"calculation": {"status": "not_provided", "expression": ""}}, []
        )
        == []
    )
    assert Decimal("23.53") * 1000 == Decimal("23530")
