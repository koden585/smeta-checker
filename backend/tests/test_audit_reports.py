"""User-visible ordering, evidence boundaries and document navigation."""

import copy
import io
from pathlib import Path
import sys
import unittest
from urllib.parse import unquote

import docx
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from audit_reports import (DISCREPANCY, MANUAL, NOT_CHECKED, build_user_report,
                           generate_excel_report, generate_word_report)


class ReportTests(unittest.TestCase):
    def make_report(self, cases):
        frame = pd.DataFrame([{"№ п.п.": i + 1, "Статус аудита": status,
                               "Детали расхождений": "found_volume=null, search_completed=true; OCR; response_status=failed",
                               "Идентификатор": "internal-id", "Тип позиции": "Работа",
                               "Формула расчета объемов работ и расхода материалов, потребности ресурсов": "10*12",
                               "Дополнительная информация (комментарий).": "Скрыто",
                               "[ВОР] Заявленный объем": "120", "[ВОР] Ед. изм.": "м³",
                               "[ПД] Запросов": 6} for i, (status, _) in enumerate(cases)], index=range(17, 17 + len(cases)))
        sources, tasks = [], []
        for status, overrides in cases:
            source = {"pdf_name": "кр.pdf", "found_page": 18, "found_pages": [18, 19],
                      "vor_pages": [18], "work_found": True, "volume_confirmed": True,
                      "unit_confirmed": True, "comment_confirmed": True,
                      "found_volume": "120", "found_unit": "м³", "evidence": [], "requests": 6}
            source.update(overrides)
            sources.append(source)
            tasks.append({"w_name": "Бетонные работы — полное наименование без обрезки на 45 символах",
                          "w_vol": "120", "w_unit": "м³", "w_comment": "", "w_page": "18",
                          "target_indices": [17], "loc_desc": "Стр. 18", "file_refs": r"C:\project\КР.pdf"})
        before_frame, before_sources = frame.copy(deep=True), copy.deepcopy(sources)
        report = build_user_report(frame, sources, tasks)
        pd.testing.assert_frame_equal(frame, before_frame)
        self.assertEqual(sources, before_sources)
        return report

    def test_priorities_stable_order_and_missing_proof_not_a_discrepancy(self):
        report = self.make_report([
            ("ОК", {}),
            ("Неполная проверка", {"found_volume": "-", "volume_confirmed": False,
                "checks": [{"claim_id": "quantity", "kind": "quantity",
                            "state": "not_found", "reasons": ["no_proposal"],
                            "required": True}]}),
            ("Ошибка", {"found_volume": "95", "volume_confirmed": False,
                "checks": [{"claim_id": "quantity", "kind": "quantity",
                            "state": "contradiction", "reasons": ["value_mismatch"],
                            "required": True}]}),
            ("Технический сбой", {"work_found": False, "found_pages": [], "found_page": None}),
            ("Ошибка", {"work_found": False, "found_pages": [], "found_page": None,
                        "workbook_notes": [
                            {"kind": "display-operand-mismatch", "display_cell": "E32",
                             "message": "в тексте расчёта E32 показано число 992"}]}),
            ("ОК", {})])
        self.assertEqual(report["№"].tolist(), [3, 5, 2, 4, 1, 6])
        self.assertEqual(report["Статус аудита"].tolist(), [DISCREPANCY, DISCREPANCY, MANUAL, NOT_CHECKED, "ОК", "ОК"])
        mismatch = report.iloc[0]
        self.assertIn("120 м³", mismatch["Что не так"])
        self.assertIn("95 м³", mismatch["Что не так"])
        self.assertIn("25 м³", mismatch["Что не так"])
        self.assertEqual(mismatch["Наименование файла"], "КР.pdf")
        self.assertEqual(mismatch["Где проверить"], "стр. 18, 19")
        for text in report["Что не так"]:
            self.assertNotIn("found_volume", text)
            self.assertNotIn("OCR", text)
        self.assertNotIn("Идентификатор", report.columns)
        contradiction = report.loc[report["№"] == 5].iloc[0]
        self.assertIn("E32", contradiction["Что не так"])
        self.assertIn("992", contradiction["Что не так"])

    def test_reference_and_unit_discrepancies_not_hidden_by_incomplete_audit(self):
        report = self.make_report([
            ("Ошибка", {"found_page": 20,
                "checks": [{"claim_id": "quantity", "kind": "quantity",
                            "state": "contradiction", "reasons": ["value_mismatch"],
                            "required": True}]}),
            ("Неполная проверка", {"unit_confirmed": False, "found_unit": "м²",
                "checks": [{"claim_id": "unit", "kind": "unit",
                            "state": "review_required",
                            "reasons": ["unit_not_bound"], "required": True}]}),
            ("Неполная проверка", {"comment_confirmed": False}),
            ("Ошибка", {"volume_confirmed": False, "found_volume": "120 м³",
                "checks": [{"claim_id": "quantity", "kind": "quantity",
                            "state": "review_required",
                            "reasons": ["ambiguous_measurements"],
                            "required": True}]})])
        self.assertEqual(report["Статус аудита"].tolist(), [DISCREPANCY, MANUAL, MANUAL, MANUAL])
        by_num = {num: text for num, text in zip(report["№"], report["Что не так"])}
        self.assertIn("не на тех страницах", by_num[1])
        self.assertIn("неполная", by_num[2])
        self.assertNotIn("Разница", by_num[4])

    def test_ok_notes_do_not_hide_actual_quantity_discrepancy(self):
        report = self.make_report([
            ("ОК", {"comment_confirmed": False, "found_page": 19}),
            ("Ошибка", {"comment_confirmed": False, "found_volume": "95",
                        "volume_confirmed": False,
                        "checks": [{"claim_id": "quantity", "kind": "quantity",
                                    "state": "contradiction",
                                    "reasons": ["value_mismatch"],
                                    "required": True}]})])
        mismatch, advisory = report.iloc[0], report.iloc[1]
        self.assertEqual(mismatch["Статус аудита"], DISCREPANCY)
        self.assertEqual(advisory["Статус аудита"], "ОК")
        self.assertIn("Примечание из ВОР", advisory["Что не так"])
        self.assertIn("рядом", advisory["Что не так"])
        self.assertNotIn("Что сделать", report.columns)

    def test_derived_volume_is_not_presented_as_printed_quantity(self):
        report = self.make_report([("Ошибка", {
            "volume_confirmed": False, "found_volume": "95",
            "checks": [{"claim_id": "quantity", "kind": "quantity",
                        "state": "contradiction", "reasons": ["value_mismatch"],
                        "required": True}],
            "volume_calculation": {"area": {"value": "950", "unit": "м2"},
                                   "thickness": {"value": "0,1", "unit": "м"}}})])
        description = report.iloc[0]["Что не так"]
        self.assertIn("по расчёту", description)
        self.assertIn("25 м³", description)

    def test_excel_colors_literal_names_and_word_navigation(self):
        report = self.make_report([
            ("ОК", {}),
            ("Ошибка", {"found_volume": "95", "volume_confirmed": False,
                "checks": [{"claim_id": "quantity", "kind": "quantity",
                            "state": "contradiction", "reasons": ["value_mismatch"],
                            "required": True}]}),
            ("Неполная проверка", {"volume_confirmed": False, "found_volume": "-",
                "checks": [{"claim_id": "quantity", "kind": "quantity",
                            "state": "not_found", "reasons": ["no_proposal"],
                            "required": True}]}),
            ("Технический сбой", {"work_found": False, "found_pages": [], "found_page": None})])
        report.loc[0, "Наименование работ"] = "=1+1"
        excel = openpyxl.load_workbook(io.BytesIO(generate_excel_report(report)))
        sheet = excel.active
        self.assertEqual([sheet.cell(i, 3).fill.fgColor.rgb for i in range(2, 6)],
                         ["00FEE2E2", "00FEF3C7", "00E5E7EB", "00DCFCE7"])
        self.assertEqual(sheet.cell(2, 2).data_type, "s")
        self.assertEqual(unquote(sheet.cell(2, 5).hyperlink.target), "КР.pdf#page=18")
        word = docx.Document(io.BytesIO(generate_word_report(report, {})))
        table = word.tables[0]
        self.assertEqual(len(table.columns), 4)
        self.assertEqual([r.cells[0].text for r in table.rows[1:]], ["2", "3", "4", "1"])
        self.assertIn("120 м³", table.rows[1].cells[3].text)
        self.assertIn("стр. 18, 19", table.rows[1].cells[3].text)
        self.assertEqual(table.rows[1].cells[1].text, "=1+1")
        targets = [unquote(rel.target_ref) for rel in word.part.rels.values() if rel.reltype.endswith("/hyperlink")]
        self.assertIn("КР.pdf#page=18", targets)
        self.assertNotIn("found_volume", " ".join(c.text for r in table.rows for c in r.cells))

    def test_missing_or_multiple_documents_have_no_invented_source_link(self):
        frame = pd.DataFrame([{"№": 1, "Статус аудита": "Ошибка"}])
        task = {"w_name": "Работа", "file_refs": "/storage/КР.pdf;..\\АР.pdf", "bind_error": "internal error"}
        report = build_user_report(frame, [{"pdf_name": "", "found_pages": []}], [task])
        self.assertEqual(report.iloc[0]["Наименование файла"], "КР.pdf\nАР.pdf")
        self.assertEqual(report.iloc[0]["Статус аудита"], NOT_CHECKED)
        sheet = openpyxl.load_workbook(io.BytesIO(generate_excel_report(report))).active
        self.assertIsNone(sheet.cell(2, 5).hyperlink)
        self.assertEqual(report.iloc[0]["Где проверить"], "Место не установлено")


class RemarkTests(unittest.TestCase):
    def _remark(self, status, source_overrides, task_overrides=None):
        frame = pd.DataFrame([{"№ п.п.": 1, "Статус аудита": status,
                               "Детали расхождений": "",
                               "[ВОР] Заявленный объем": "120", "[ВОР] Ед. изм.": "м³"}])
        source = {"pdf_name": "кр.pdf", "found_page": 18, "found_pages": [18],
                  "vor_pages": [18], "work_found": True, "volume_confirmed": True,
                  "unit_confirmed": True, "comment_confirmed": True,
                  "found_volume": "120", "found_unit": "м³", "evidence": []}
        source.update(source_overrides)
        task = {"w_name": "Бетон", "w_vol": "120", "w_unit": "м³", "w_comment": "",
                "w_page": "18", "target_indices": [17], "loc_desc": "Стр. 18",
                "file_refs": r"C:\project\КР.pdf"}
        task.update(task_overrides or {})
        report = build_user_report(frame, [source], [task])
        self.assertNotIn("Что сделать", report.columns)
        return report.iloc[0]

    def test_quantity_mismatch_names_numbers_and_place(self):
        row = self._remark("Ошибка", {"found_volume": "95", "volume_confirmed": False,
            "checks": [{"claim_id": "quantity", "kind": "quantity",
                        "state": "contradiction", "reasons": ["value_mismatch"],
                        "required": True}]})
        self.assertEqual(row["Статус аудита"], DISCREPANCY)
        self.assertIn("120", row["Что не так"])
        self.assertIn("95", row["Что не так"])
        self.assertIn("стр. 18", row["Где проверить"])

    def test_equal_unconfirmed_pair_is_not_reported_as_missing(self):
        row = self._remark("Неполная проверка", {"volume_confirmed": False,
            "found_volume": "1879", "found_unit": "м²",
            "checks": [{"claim_id": "quantity", "kind": "quantity",
                        "state": "review_required",
                        "reasons": ["ambiguous_measurements"], "required": True}]},
            {"w_vol": "1879.0", "w_unit": "м2"})
        self.assertEqual(row["Статус аудита"], MANUAL)
        self.assertNotIn("не нашёлся", row["Что не так"])
        self.assertIn("1879", row["Что не так"])

    def test_technology_difference_names_both_methods(self):
        row = self._remark("Ошибка", {"volume_confirmed": False, "found_volume": "-",
            "technology": {"status": "different",
            "expected": "втулки", "actual": "стяжки"}})
        self.assertEqual(row["Статус аудита"], DISCREPANCY)
        self.assertIn("втулки", row["Что не так"])
        self.assertIn("стяжки", row["Что не так"])

    def test_technology_difference_with_confirmed_quantity_is_manual(self):
        row = self._remark("Неполная проверка", {"technology": {"status": "different",
            "expected": "втулки", "actual": "стяжки"}})
        self.assertEqual(row["Статус аудита"], MANUAL)
        self.assertIn("инженерная сверка", row["Что не так"])

    def test_unconfirmed_technology_names_cross_pages(self):
        row = self._remark("Неполная проверка", {"technology": {
            "status": "unconfirmed", "expected": "укладка",
            "actual": "укрепление", "cross_pages": [34, 35]}})
        self.assertEqual(row["Статус аудита"], MANUAL)
        self.assertIn("стр. 34, 35", row["Что не так"])
    def test_consumption_mismatch_names_delta_and_derived_rate(self):
        row = self._remark("Ошибка", {"volume_confirmed": False, "found_volume": "-",
            "consumption": {"matches": False, "schedule_volume": "86,79",
                "schedule_page": 88, "schedule_item": "Щебень",
                "delta": "-7,89", "delta_percent": "-9,09",
                "rate": "0,15", "unit": "м3/м2", "derived_rate": "0,165",
                "uncertainties": []}})
        self.assertEqual(row["Статус аудита"], DISCREPANCY)
        self.assertIn("86,79", row["Что не так"])
        self.assertIn("0,165", row["Что не так"])
        self.assertIn("88", row["Где проверить"])

    def test_incomplete_says_what_is_missing_and_where(self):
        row = self._remark("Неполная проверка", {"volume_confirmed": False, "found_volume": "-"})
        self.assertEqual(row["Статус аудита"], MANUAL)
        self.assertIn("стр. 18", row["Где проверить"])

    def test_technical_failure_has_no_invented_pages(self):
        row = self._remark("Технический сбой", {"work_found": False, "found_pages": [],
                                                "found_page": None, "vor_pages": []})
        self.assertEqual(row["Статус аудита"], NOT_CHECKED)
        self.assertEqual(row["Где проверить"], "Место не установлено")

    def test_ok_with_note_stays_ok(self):
        row = self._remark("ОК", {"comment_confirmed": False})
        self.assertEqual(row["Статус аудита"], "ОК")

    def test_display_mismatch_is_discrepancy(self):
        row = self._remark("ОК", {"workbook_notes": [
            {"kind": "display-operand-mismatch", "display_cell": "E20",
             "message": "в тексте расчёта E20 показано число 631.2"}]})
        self.assertEqual(row["Статус аудита"], DISCREPANCY)
        self.assertIn("В самой ведомости противоречие", row["Что не так"])
        self.assertIn("631.2", row["Что не так"])

    def test_scheme_question_stays_ok(self):
        row = self._remark("ОК", {"workbook_notes": [
            {"kind": "scheme-divisor-mismatch", "cell": "M35",
             "message": "позиция №17 считает по схеме M35 = 4"}]})
        self.assertEqual(row["Статус аудита"], "ОК")
        self.assertIn("Вопрос к исходным данным", row["Что не так"])

    def test_chain_convergence_is_note_not_mismatch(self):
        row = self._remark("ОК", {"consumption": {
            "matches": True, "schedule_volume": "86,79", "schedule_page": 88,
            "schedule_item": "Щебень", "rate": "0,15", "unit": "м3/м2",
            "uncertainties": ["rate-justification-open"],
            "stock_chain": {"material_number": "11", "factor": "1.1",
                            "work_volume": "78,90",
                            "material_volume": "86,79"}}})
        self.assertEqual(row["Статус аудита"], "ОК")
        self.assertIn("Потребность в материале сходится", row["Что не так"])
        self.assertIn("× 1.1", row["Что не так"])
        self.assertIn("Обоснованность нормы", row["Что не так"])

    def test_chain_mismatch_names_work_factor_material(self):
        row = self._remark("Ошибка", {"volume_confirmed": False,
            "found_volume": "-",
            "consumption": {"matches": False, "schedule_volume": "80,00",
                "schedule_page": 88, "schedule_item": "Щебень",
                "delta": "6,79", "delta_percent": "8,49",
                "rate": "0,15", "unit": "м3/м2", "uncertainties": [],
                "stock_chain": {"material_number": "11", "factor": "1.1",
                                "work_volume": "78,90",
                                "material_volume": "86,79"}}})
        self.assertEqual(row["Статус аудита"], DISCREPANCY)
        self.assertIn("78,90", row["Что не так"])
        self.assertIn("× 1.1", row["Что не так"])
        self.assertIn("№11", row["Что не так"])

    def test_found_off_reference_with_matching_quantity_is_manual(self):
        row = self._remark("Неполная проверка", {"found_page": 20,
                                                "vor_pages": [18]})
        self.assertEqual(row["Статус аудита"], MANUAL)
        self.assertIn("не на тех страницах", row["Что не так"])

    def test_missing_reference_link_is_manual_not_contradiction(self):
        row = self._remark("Неполная проверка", {},
                           {"target_indices": [], "loc_desc": "Страница не указана"})
        self.assertEqual(row["Статус аудита"], MANUAL)
        self.assertIn("нет годной ссылки", row["Что не так"])

if __name__ == "__main__":
    unittest.main()
