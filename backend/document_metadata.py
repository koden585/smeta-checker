"""Bounded, deterministic document headers; no OCR or model calls."""

import csv
import io
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

EXTRACTOR_VERSION = "1"
_FIELDS = ("construction_title", "raw_title", "document_code", "pd_reference")
_VOR_LABELS = {
    "наименование стройки": "construction_title",
    "стройка": "construction_title",
    "наименование объекта капитального строительства": "raw_title",
    "объект капитального строительства": "raw_title",
    "наименование объекта": "raw_title",
    "объект строительства": "raw_title",
    "объект": "raw_title",
    "наименование ведомости": "raw_title",
    "наименование документа": "raw_title",
    "название документа": "raw_title",
    "ведомость объемов работ": "document_code",
    "номер ведомости": "document_code",
    "номер документа": "document_code",
    "обозначение документа": "document_code",
    "код документа": "document_code",
    "шифр проектной документации": "pd_reference",
    "шифр проекта": "pd_reference",
    "проектная документация": "pd_reference",
    "шифр пд": "pd_reference",
    "ссылка на пд": "pd_reference",
    "основание": "pd_reference",
    "шифр": "document_code",
}
_PDF_LABELS = {
    **_VOR_LABELS,
    "наименование чертежа": "raw_title",
    "наименование комплекта": "raw_title",
    "шифр проектной документации": "document_code",
    "шифр проекта": "document_code",
    "обозначение": "document_code",
    "шифр документа": "document_code",
    "наименование объекта капитального строительства": "construction_title",
    "объект капитального строительства": "construction_title",
    "наименование объекта": "construction_title",
    "объект строительства": "construction_title",
    "объект": "construction_title",
}


def _flat(value) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def _label(text: str, labels: dict) -> tuple[str, str] | None:
    folded = text.casefold().replace("ё", "е")
    for name in sorted(labels, key=len, reverse=True):
        if folded.startswith(name):
            rest = text[len(name) :]
            if not rest or not rest[0].isalpha():
                return labels[name], rest.lstrip(" :№–—-.")
    return None


def _table_header(values: list[str]) -> bool:
    folded = [value.casefold() for value in values]
    return any(re.match(r"№\s*п\.?\s*п\.?", value) for value in folded) or (
        any("наименование работ" in value for value in folded)
        and any("ед. изм" in value or "количество" in value for value in folded)
    )


def _put(
    result: dict,
    conflicts: set,
    field: str,
    value: str,
    sheet: str,
    cell: str,
    raw: str,
):
    value = _flat(value)
    if not value or len(value) > 2048 or not any(c.isalnum() for c in value):
        return
    if field in conflicts:
        return
    old = result[field]
    if old and old != value:
        # Two contradictory declarations cannot establish a document identity.
        conflicts.add(field)
        result[field] = ""
        result["sources"].pop(field, None)
    elif not old:
        result[field] = value
        result["sources"][field] = {"sheet": sheet, "cell": cell, "raw": raw}


def _scan_rows(rows: list[list], sheet: str, result: dict, conflicts: set):
    for index, row in enumerate(rows):
        values = [_flat(value) for value in row]
        if _table_header(values):
            break
        for col, text in enumerate(values):
            hit = _label(text, _VOR_LABELS)
            if hit is None:
                continue
            field, value = hit
            source_row, source_col, raw = index, col, str(row[col] or "")
            if not value:
                following = next(
                    ((c, v) for c, v in enumerate(values[col + 1 :], col + 1) if v),
                    None,
                )
                if following and _label(following[1], _VOR_LABELS) is None:
                    source_col, value = following
                    raw = str(row[source_col])
                elif index + 1 < len(rows) and not _table_header(
                    [_flat(v) for v in rows[index + 1]]
                ):
                    below = rows[index + 1][col] if col < len(rows[index + 1]) else None
                    candidate = _flat(below)
                    if candidate and _label(candidate, _VOR_LABELS) is None:
                        source_row, value, raw = index + 1, candidate, str(below)
            _put(
                result,
                conflicts,
                field,
                value,
                sheet,
                f"{get_column_letter(source_col + 1)}{source_row + 1}",
                raw,
            )


def _scan_pdf(data: bytes, result: dict, conflicts: set):
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as document:
        for page_index in range(min(3, document.page_count)):
            lines = document[page_index].get_text("text").splitlines()[:300]
            for index, line in enumerate(lines):
                hit = _label(_flat(line), _PDF_LABELS)
                if hit is None:
                    continue
                field, value = hit
                source_index = index
                if not value and index + 1 < len(lines):
                    candidate = _flat(lines[index + 1])
                    if _label(candidate, _PDF_LABELS) is None:
                        source_index, value = index + 1, candidate
                _put(
                    result,
                    conflicts,
                    field,
                    value,
                    f"page {page_index + 1}",
                    f"L{source_index + 1}",
                    lines[source_index],
                )


def extract_document_metadata(data: bytes, filename: str, kind: str) -> dict:
    """Keep extracted text and cell provenance separate from editable overrides.

    XLSX scans at most 100 preamble rows and 64 columns per sheet. Merged
    cells naturally expose their value only at the top-left source address.
    Unlabelled cells never establish identity. Contradictory values stay unknown.
    """
    result = {field: "" for field in _FIELDS}
    result.update(sources={}, extractor_version=EXTRACTOR_VERSION)
    conflicts: set[str] = set()
    try:
        if kind == "pdf":
            _scan_pdf(data, result, conflicts)
        elif filename.lower().endswith(".csv"):
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1251")
            # Inspect only a bounded preamble, not the full work table.
            preamble = "\n".join(text.splitlines()[:100])
            try:
                dialect = csv.Sniffer().sniff(preamble, delimiters=";\t,")
            except csv.Error:
                dialect = csv.excel_tab if "\t" in preamble else csv.excel
                if dialect is csv.excel:
                    rows = list(csv.reader(io.StringIO(preamble), delimiter=";"))
                else:
                    rows = list(csv.reader(io.StringIO(preamble), dialect))
            else:
                rows = list(csv.reader(io.StringIO(preamble), dialect))
            _scan_rows([row[:64] for row in rows], "", result, conflicts)
        else:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                for sheet in book.worksheets:
                    rows = [
                        list(row)
                        for row in sheet.iter_rows(
                            max_row=min(sheet.max_row or 100, 100),
                            max_col=min(sheet.max_column or 64, 64),
                            values_only=True,
                        )
                    ]
                    _scan_rows(rows, sheet.title, result, conflicts)
            finally:
                book.close()
    except ValueError, OSError, KeyError, csv.Error, BadZipFile, RuntimeError:
        # Import validates file contents separately; metadata is optional.
        return {
            **{field: "" for field in _FIELDS},
            "sources": {},
            "extractor_version": EXTRACTOR_VERSION,
        }
    return result
