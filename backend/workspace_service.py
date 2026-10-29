"""Read persisted audit inputs/results. No OCR or model requests on reads."""

from __future__ import annotations

import io
from collections import Counter
from html.parser import HTMLParser
from html import escape
from uuid import UUID
from typing import Literal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from audit_models import VorPosition
from workspace_models import ResultDTO, RunDTO

def run_progress(context, record) -> RunDTO:
    if "job_counts" not in record:
        record = context.repository.run_progress_records(run_id=record["id"])[0]
    counts = Counter(record["job_counts"])
    job_total = sum(counts.values())
    total = (
        record["input_total"]
        if record["input_total"] is not None
        else job_total or None
    )
    status = record["status"]
    finished = counts["done"] + counts["failed"]
    stage = (
        status
        if status in {"done", "failed"}
        else (
            "preparing"
            if record["prepare_seconds"] is None
            and (not job_total or record["pdf_jobs"])
            else "finalizing"
            if total is not None and finished == total
            else "auditing"
            if job_total
            else "preparing"
        )
    )
    return RunDTO(
        run_id=str(record["id"]),
        project_id=str(record["project_id"]),
        vor_version_id=str(record["vor_version_id"]),
        vor_filename=record["vor_filename"],
        status=status,
        stage=stage,
        total=total,
        done=counts["done"] if job_total else record["result_count"],
        failed=counts["failed"],
        running=counts["running"],
        queued=counts["queued"],
        created_at=record.get("created_at"),
        finished_at=record.get("finished_at"),
        preparation_done=(record.get("pdf_preparation") or {}).get("done", 0),
        preparation_total=(record.get("pdf_preparation") or {}).get("total"),
        normative_counts=record.get("normative_counts") or {},
    )


def inputs_for_run(context, record) -> list[VorPosition]:
    snapshot = record["config"].get("position_inputs")
    if isinstance(snapshot, list):
        return [VorPosition.model_validate(item) for item in snapshot]
    from vor_parse import load_vor_positions

    version = context.repository.load_version(record["vor_version_id"])
    positions, _ = load_vor_positions(
        context.blob_store.get(version.blob_key), version.filename
    )
    return positions


def _provenance(position: VorPosition) -> tuple[str, int | None, dict[str, str]]:
    source = position.position_source
    if source is None:
        return "", None, {}
    return source.sheet, source.row, dict(source.cells)


def present_results(context, record) -> list[ResultDTO]:
    from normative_local_checks import LOCAL_RULES_VERSION
    context.repository.schedule_local_normative_refresh(record['id'], LOCAL_RULES_VERSION)
    positions = {p.index: p for p in inputs_for_run(context, record)}
    from vor_material_links import material_links
    links = material_links([p.legacy_task() for p in positions.values()])
    validation = context.repository.normative_validation_context(record['id'])
    jobs = {
        j.position_index: j for j in context.repository.list_position_jobs(record["id"])
    }
    result = []
    for item in record["indexed_results"]:
        index, stored = item["position_index"], item["result"]
        position = positions.get(index)
        if position is None:
            raise ValueError("Исходная позиция результата не найдена")
        job = jobs.get(index)
        if record["status"] == "running" and job is not None and job.status not in {"done", "failed"}:
            continue
        payload = job.payload if job else {}
        task = position.legacy_task()
        task["derived_material"] = links.get(index)
        task.update(
            {
                "bind_error": task.get("bind_error") or payload.get("bind_error"),
                "target_indices": payload.get(
                    "target_indices", [p - 1 for p in stored.get("vor_pages", [])]
                ),
                "loc_desc": task.get("loc_desc") or payload.get("loc_desc", ""),
            }
        )
        from audit_review_policy import review_source
        stored = review_source(stored, task)
        tracking = {
            **stored,
            "found_page": stored.get("found_page")
            or next(iter(stored.get("found_pages") or []), None),
        }
        from audit_reports import build_user_report

        report = build_user_report(
            [{"№": position.number, "Статус аудита": stored["status"]}],
            [tracking],
            [task],
        )[0]
        matching_report_status = report["Статус аудита"]
        from normative_verdict import publish, verified_quantity
        from human_basis import position_key
        report = publish(report, stored, position_key(position), validation.get(index),
            enabled=bool(record['config'].get('normative_after_match')), derived=bool(links.get(index)))
        report_status = report["Статус аудита"]
        quantity = verified_quantity((validation.get(index) or {}).get('review'),
            position_key(position), position.unit, derived=bool(links.get(index)))
        queues: dict[str, Literal['fix', 'review', 'missing', 'ok', 'pending']] = {
            "Обнаружено расхождение": "fix",
            "Требуется ручная проверка": "review",
            "Проверка не выполнена": "missing",
            "ОК": "ok",
            "Проверяется по нормативам": "pending",
        }
        queue = queues[report_status]
        sheet, source_row, cells = _provenance(position)
        where = str(report.get("Где проверить") or "")
        filename = str(report.get("Наименование файла") or "")
        if filename and filename != "Не указан":
            action = f"{filename}: {where}" if where else filename
        else:
            action = where or "Откройте позицию и проверьте источники вручную."
        result.append(
            ResultDTO(
                id=f"{record['id']}:{index}",
                position_index=index,
                number=position.number,
                name=position.work_name,
                unit=position.unit,
                declared_quantity=position.volume,
                found_quantity=str(quantity[0]) if quantity else str(stored["found_volume"])
                if stored.get("found_volume") not in {None, "", "-"}
                else None,
                found_unit=quantity[1] if quantity else stored.get("found_unit")
                if stored.get("found_unit") not in {None, "", "-"}
                else None,
                formula=position.formula,
                sheet=sheet,
                source_row=source_row,
                cells=cells,
                audit_status=stored["status"],
                report_status=report_status,
                matching_report_status=matching_report_status,
                related_work=links.get(index),
                queue=queue,
                problem=report["Что не так"],
                action=action,
                work_found=bool(stored.get("work_found")),
                volume_confirmed=bool(stored.get("volume_confirmed")) or bool(quantity),
                unit_confirmed=bool(stored.get("unit_confirmed")) or bool(quantity),
                comment_confirmed=bool(stored.get("comment_confirmed")),
                details=report['Что не так'] if record['config'].get('normative_after_match') or validation.get(index)
                else stored.get("discrepancy_details") or "",
            )
        )
    from normative_verdict import apply_parent_verdicts
    return apply_parent_verdicts(result, enabled=bool(record['config'].get('normative_after_match')) or bool(validation),
        running=record['status'] == 'running')


def validated_word_report(context, record):
    """Export the same final verdict as the UI; never serve stale red conclusions."""
    validation = context.repository.normative_validation_context(record['id'])
    if not record['config'].get('normative_after_match') and not validation:
        return None
    rows = present_results(context, record)
    if any(r.queue == 'pending' for r in rows):
        raise ValueError('Нормативная валидация ещё выполняется. Итоговый отчёт будет доступен после её завершения.')
    from audit_reports import generate_word_report
    report = [{'№': r.number, 'Наименование работ': r.name, 'Статус аудита': str(r.report_status),
        'Что не так': r.problem, 'Наименование файла': 'Не указан', 'Где проверить': r.action} for r in rows]
    return generate_word_report(report, {'object_name': record['config'].get('project') or record['vor_filename']})


def sheet_window(data: bytes, position: VorPosition):
    sheet = position.position_source.sheet if position.position_source else ""
    source_row = position.position_source.row if position.position_source else None
    if not sheet or source_row is None:
        return None
    book = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    try:
        sheet_obj = book[sheet]
        low = high = source_row
        # One reviewed row only. Identifiers and type remain internal metadata.
        columns = list(range(1, 6))
        merged = list(sheet_obj.merged_cells.ranges)
        coords = (
            set((position.position_source.cells or {}).keys())
            if position.position_source
            else set()
        )
        rows = []
        for number in range(low, high + 1):
            cells = []
            for col in columns:
                cell = sheet_obj.cell(number, col)
                span = next(
                    (
                        m
                        for m in merged
                        if m.min_row <= number <= m.max_row
                        and m.min_col <= col <= m.max_col
                    ),
                    None,
                )
                if span and col != span.min_col:
                    continue
                if span:
                    cell = sheet_obj.cell(span.min_row, span.min_col)
                cells.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": str(cell.value) if cell.value is not None else "",
                        "column_span": min(span.max_col, columns[-1]) - col + 1 if span else 1,
                        "row_span": 1,
                        "highlight": cell.coordinate in coords and number == source_row,
                    }
                )
            rows.append(
                {
                    "number": number,
                    "selected": number == source_row,
                    "cells": cells,
                }
            )
        return {
            "sheet": sheet_obj.title,
            "columns": [get_column_letter(c) for c in columns],
            "rows": rows,
        }
    finally:
        book.close()


class _CellText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.cells: list[str] = []
        self._current: list[str] = []
        self._in_cell = False

    def handle_starttag(self, tag, attrs):
        if tag in {"td", "th"}:
            self._current = []
            self._in_cell = True

    def handle_endtag(self, tag):
        if tag in {"td", "th"} and self._in_cell:
            text = "".join(self._current).strip()
            if text:
                self.cells.append(text)
            self._in_cell = False

    def handle_data(self, data):
        if self._in_cell:
            self._current.append(data)


def _quote_boxes(page, quote):
    """Locate a quote on a PDF page, tolerating table HTML in evidence.

    Evidence quotes may be whole ``<tr>`` rows from the table layer; raw
    tags never occur in the page text, so split into cells and match the
    row whose cells share one text row: the longest cell via its longest
    unique fragment, numeric cells by row intersection. Plain-text quotes
    keep the previous exact behaviour.
    """
    import pymupdf

    normalized = " ".join(page.get_text().casefold().split())
    if "<td" not in quote and "<th" not in quote:
        phrase = " ".join(quote.casefold().split())
        if not phrase or normalized.count(phrase) != 1:
            return []
        return list(page.search_for(quote))
    parser = _CellText()
    try:
        parser.feed(quote)
    except Exception:
        return []
    phrases = [" ".join(cell.split()) for cell in parser.cells]
    phrases = [phrase for phrase in phrases if phrase]
    if not phrases:
        return []
    texts = [phrase for phrase in phrases if len(phrase) >= 4]
    numbers = {token for phrase in phrases for token in phrase.split()
               if token.replace(",", "").replace(".", "").isdigit()}
    words = page.get_text("words")
    boxes = []
    for phrase in texts:
        unique = _unique_subphrase(normalized, phrase)
        if unique:
            boxes.extend(page.search_for(unique))
    if boxes:
        return [box for box in boxes if box.width > 0 and box.height > 0]
    # Repeated work name: disambiguate by the numeric cells sharing its row.
    longest = max(texts, key=len) if texts else ""
    if not longest or not numbers:
        return []
    candidates = page.search_for(longest)
    rows = []
    for candidate in candidates:
        band = pymupdf.Rect(0, candidate.y0 - candidate.height,
                         page.rect.width, candidate.y1 + candidate.height)
        row_words = [word for word in words
                     if band.intersects(pymupdf.Rect(word[:4]))]
        row_text = " ".join(word[4] for word in row_words)
        if any(number in row_text for number in numbers):
            rows.extend(box for box in page.search_for(longest)
                        if band.intersects(box))
    return [box for box in rows if box.width > 0 and box.height > 0]


def _unique_subphrase(normalized, phrase):
    """Longest leading word-run occurring exactly once (min 4 chars).

    Table rows repeat across pages (headers, duplicated lines): the full
    cell text may occur twice while a leading fragment is unique. Never
    return an ambiguous or trivially short phrase.
    """
    words = str(phrase or "").split()
    for end in range(len(words), 0, -1):
        candidate = " ".join(words[:end])
        if len(candidate) < 4:
            break
        if normalized.count(candidate.casefold()) == 1:
            return candidate
    return ""


def _addressed_evidence(page, evidence, stored, position):
    """Locate the observed quantity's physical row, including historical results.

    This only presents already evaluated evidence. The VOR expected quantity
    never chooses a row here; ambiguous or unconfirmed observations stay open.
    """
    import re
    import pymupdf
    from audit_core import _bind_quantity_fact, _compare_bound_fact, _shares_work_object, _terms

    quote = str(evidence.get("quote") or "")
    actual = str(stored.get("found_volume") or "")
    unit = str(stored.get("found_unit") or "")
    if not stored.get("volume_confirmed") or actual in {"", "-"}:
        return quote, []
    table_match = re.search(r'\bdata-table="([^"]+)"', quote)
    locators = []
    if evidence.get("locator"):
        locators.append(evidence["locator"])
    for check in stored.get("checks") or []:
        for fact in check.get("facts") or []:
            locators.extend(fact.get("locators") or [])
    candidates = {}
    query = _terms(position.work_name)
    for locator in locators:
        if locator.get("page") != evidence.get("page") or locator.get("quality") != "usable":
            continue
        if table_match and locator.get("table_id") != table_match.group(1):
            continue
        rows = {}
        for cell in locator.get("members") or []:
            if cell.get("role") in {"header", "group"} or cell.get("quality") != "usable":
                continue
            rows.setdefault(cell.get("row_id"), []).append(cell)
        for row_id, cells in rows.items():
            if not row_id:
                continue
            cells.sort(key=lambda cell: cell.get("column", 0))
            texts = [" ".join(cell.get("text", "").split()) for cell in cells]
            row_quote = "<tr>" + "".join("<td>" + escape(text) + "</td>" for text in texts) + "</tr>"
            if row_quote not in quote:
                continue
            text = " | ".join(texts)
            if not _shares_work_object(text, position.work_name):
                continue
            fact, bound = _bind_quantity_fact([text], position.work_name, unit)
            if fact is None or bound != "bound" or _compare_bound_fact(fact, actual, unit) != "equal":
                continue
            boxes = []
            for cell in cells:
                box = cell.get("bbox") or []
                if len(box) != 4:
                    break
                rect = pymupdf.Rect(box)
                if rect.is_empty or rect.is_infinite or not page.cropbox.contains(rect):
                    break
                boxes.append(rect)
            else:
                score = len(query & _terms(text))
                candidates[(locator.get("table_id"), row_id)] = (score, row_quote, boxes)
    if not candidates:
        return quote, []
    best = max(item[0] for item in candidates.values())
    selected = [item for item in candidates.values() if item[0] == best]
    if len(selected) != 1:
        return quote, []
    _, row_quote, boxes = selected[0]
    return row_quote, boxes



def source_evidence(context, record, index, position):
    import pymupdf

    stored = next(
        x["result"] for x in record["indexed_results"] if x["position_index"] == index
    )
    versions = record["config"].get("pdf_versions") or {}
    result = []
    for evidence in stored.get("evidence") or []:
        name = evidence.get("document_id") or stored.get("pdf_name") or ""
        version_id = versions.get(name) or (
            stored.get("pdf_version_id") if name == stored.get("pdf_name") else None
        )
        if not version_id or not isinstance(evidence.get("page"), int):
            continue
        version = context.repository.load_version(UUID(version_id))
        if context.repository.version_project_id(version.id) != record["project_id"]:
            continue
        page_number, quote = evidence["page"], str(evidence.get("quote") or "")
        with pymupdf.open(
            stream=context.blob_store.get(version.blob_key), filetype="pdf"
        ) as pdf:
            if not 1 <= page_number <= len(pdf):
                continue
            page = pdf[page_number - 1]
            quote, boxes = _addressed_evidence(page, evidence, stored, position)
            if not boxes:
                boxes = _quote_boxes(page, quote)
            # Repeated or long/multiline ambiguous hits do not become an accepted highlight.
            union = pymupdf.Rect(boxes[0]) if boxes else None
            for box in boxes[1:]:
                union |= box
            if union and union.height > page.rect.height * 0.2:
                boxes = []
            region = {
                "width": page.rect.width,
                "height": page.rect.height,
                "highlights": [
                    {
                        "bbox": list(b * page.rotation_matrix),
                        "kind": "quote",
                        "text": quote,
                    }
                    for b in boxes
                ],
                "reason": "Область основания найдена на странице"
                if boxes
                else "Координаты цитаты не установлены; проверьте страницу вручную",
            }
        result.append(
            {
                "version_id": version_id,
                "filename": version.filename,
                "page": page_number,
                "quote": quote,
                "kind": evidence.get("kind", "work"),
                "region": region,
                "confirmed": True,
            }
        )
    if not result:
        # Navigation to the VOR reference never constitutes proof.
        name = stored.get("pdf_name") or position.document_id
        version_id = versions.get(name)
        pages = stored.get("vor_pages") or position.primary_pages
        if version_id and pages:
            version = context.repository.load_version(UUID(version_id))
            if (
                context.repository.version_project_id(version.id)
                == record["project_id"]
            ):
                result.append(
                    {
                        "version_id": version_id,
                        "filename": version.filename,
                        "page": pages[0],
                        "quote": "Ссылка из ВОР; подтверждение не установлено",
                        "kind": "reference",
                        "region": None,
                        "confirmed": False,
                    }
                )
    return result
