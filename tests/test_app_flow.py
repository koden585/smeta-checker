"""Exercise the Streamlit audit and exports with a synthetic project and mocked API."""

import csv
import io
import json
import os
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import fitz
import pandas as pd
import docx
from streamlit.testing.v1 import AppTest

from audit_models import AuditSettings


def _yandex_only_settings(*args, **kwargs):
    """UI tests pin the Yandex backend: the real local .env may hold AnyModel."""
    return AuditSettings(
        yandex_cloud_folder="test-folder",
        yandex_cloud_api_key="test-only",
        allowed_llm_hosts=("ai.api.cloud.yandex.net",),
        database_url=os.environ.get("DATABASE_URL", ""),
        blob_store_url=os.environ.get("BLOB_STORE_URL", ""),
    )


class FakeClient:
    def __init__(self, **kwargs):
        self.responses = SimpleNamespace(create=self.create)

    def create(self, *, model, instructions, input, max_output_tokens, **_):
        prompt = input
        if "Проверяемая работа: Pump station" in prompt:
            raise RuntimeError("Simulated provider failure")
        import re
        sources = re.findall(r"\[([a-z0-9]+)\] ([^\n]+)", prompt)
        work_sid = next(sid for sid, text in sources if "Copper cable" in text)
        quantity_sid = next(sid for sid, text in sources if "45" in text)
        data = {
            "claims": [
                {"claim_id": "presence", "source_ids": [work_sid]},
                {"claim_id": "quantity", "source_ids": [quantity_sid],
                 "value_source_id": quantity_sid, "unit_source_id": quantity_sid},
                {"claim_id": "unit", "source_ids": [quantity_sid]},
            ],
            "discrepancy_details": "Confirmed with multi-page evidence",
        }
        return SimpleNamespace(output_text=json.dumps(data), usage=SimpleNamespace(total_tokens=100))

    def close(self):
        pass


class PipelineTests(unittest.TestCase):
    def test_first_request_overlaps_second_preparation_without_pdf_cross_thread_access(self):
        import threading
        from audit_service import run_audit
        from pdf_text import PDFTextSource
        started, prepared = threading.Event(), threading.Event()
        owner = threading.get_ident()
        preparation_threads, progress_threads = [], []
        snapshots = []
        test = self
        class Client(FakeClient):
            def create(self, **kwargs):
                if "Проверяемая работа: Copper cable" in kwargs["input"]:
                    started.set()
                    if not prepared.wait(5):
                        raise RuntimeError("Preparation did not overlap request")
                return super().create(**kwargs)
        with tempfile.TemporaryDirectory() as cache, fitz.open() as pdf:
            for content in ("Copper cable specification. " * 20, "Copper cable 45 m. " * 20):
                pdf.new_page().insert_textbox(fitz.Rect(50, 50, 500, 600), content)
            class Source(PDFTextSource):
                def __init__(self, doc):
                    super().__init__(doc, cache_dir=cache)
                def for_position(self, *args, **kwargs):
                    preparation_threads.append(threading.get_ident())
                    if snapshots:
                        test.assertTrue(started.wait(5), "First request must start before remaining preparation")
                    result = super().for_position(*args, **kwargs)
                    snapshots.append(result)
                    if len(snapshots) == 2:
                        prepared.set()
                    return result
            vor = io.BytesIO(("№ п.п.;Наименование работ;Объём работ / Количество;Ед. изм.;"
                              "Ссылка на ПД;Наименование файла;Номер страниц (через пробел)\n"
                              "1;Copper cable;45;m;;project.pdf;1\n"
                              "2;Pump station;1;шт;;project.pdf;2\n").encode("utf-8-sig"))
            with patch("audit_service.PDFTextSource", Source):
                output = run_audit({"vor.csv": vor}, ["vor.csv"], {"project.pdf": pdf},
                                   api_key="test", folder="test", model_name="test",
                                   client_factory=Client,
                                   progress=lambda *_: progress_threads.append(threading.get_ident()))
            self.assertEqual(output["issues"], [])
            table = output["results"]["vor.csv"]["df"]
            self.assertEqual(table["Статус аудита"].tolist(), ["ОК", "Технический сбой"])
            self.assertEqual(preparation_threads, [owner, owner])
            self.assertTrue(all(thread == owner for thread in progress_threads))
            self.assertIsInstance(snapshots[0].texts, tuple)


class AppFlowTests(unittest.TestCase):
    def _run_audit(self, folder, header, rows, preparation_fault_on_call=None, extra_env=None):
        with fitz.open() as pdf:
            for content in ("Copper cable specification. " * 20, "Copper cable 45 m. " * 20):
                page = pdf.new_page()
                page.insert_textbox(fitz.Rect(50, 50, 500, 600), content)
            pdf.save(folder / "project.pdf")
        with (folder / "vor.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=";")
            writer.writerow(header)
            writer.writerows(rows)
        from pdf_text import PDFTextSource
        real_for_position = PDFTextSource.for_position
        state = {"calls": 0}

        def _maybe_flaky(source, *args, **kwargs):
            state["calls"] += 1
            if preparation_fault_on_call == state["calls"]:
                raise RuntimeError("Simulated preparation failure")
            return real_for_position(source, *args, **kwargs)

        env = {"YANDEX_CLOUD_API_KEY": "test-only", "YANDEX_CLOUD_FOLDER": "test-folder"}
        env.update(extra_env or {})
        with patch.dict(os.environ, env), \
                patch("openai.OpenAI", FakeClient), \
                patch("audit_tracing._configured_client", return_value=None), \
                patch("app.AuditSettings.load", side_effect=_yandex_only_settings), \
                patch("pdf_text.PDFTextSource.for_position", _maybe_flaky):
            at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=30)
            at.run()
            self.assertEqual(len(at.exception), 0, str(at.exception))
            self.assertEqual([i for i in at.selectbox if i.label in ("Провайдер", "Модель")], [])
            folder_input = next(item for item in at.text_input if item.label == "Путь к папке проекта на диске")
            folder_input.set_value(str(folder)).run()
            self.assertEqual(len(at.exception), 0, str(at.exception))
            next(item for item in at.button if item.label == "🚀 Запустить аудит").click().run()
            self.assertEqual(len(at.exception), 0, str(at.exception))
            return at

    def test_audit_continues_after_failed_row_and_exports_actual_page(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1],
                 [2, "Pump station", 1, "шт", "", "project.pdf", 2],
                 [3, "Copper cable", 45, "m", "", "project.pdf", 2]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(table["Статус аудита"].tolist(), ["ОК", "Технический сбой", "ОК"])
            self.assertEqual(table["[ПД] Фактическая страница"].tolist(), ["Стр. 1, 2", "-", "Стр. 1, 2"])
            self.assertEqual(table["[ПД] Найденная единица"].tolist(), ["м", "-", "м"])
            self.assertEqual(table["[ПД] Найденные страницы"].tolist(), ["1, 2", "-", "1, 2"])
            self.assertTrue(all("Работа: да" in v for v in table["[ПД] Подтверждения (работа/объём/ед./прим.)"].tolist()[:1] + table["[ПД] Подтверждения (работа/объём/ед./прим.)"].tolist()[2:]))
            self.assertIn("Copper cable 45 m", table["[ПД] Доказательства (цитаты со страницами)"].iloc[0])
            self.assertIn("Copper cable 45 m", table["[ПД] Доказательства (цитаты со страницами)"].iloc[2])
            self.assertEqual(table["[ПД] Поиск завершён"].tolist(), ["да", "нет", "да"])
            tracking = saved["tracking"]
            self.assertEqual(tracking[0]["found_pages"], [1, 2])
            self.assertEqual(tracking[2]["found_pages"], [1, 2])
            self.assertEqual(tracking[0]["found_page"], 1)
            self.assertEqual(len(tracking[0]["evidence"]), 3)
            self.assertEqual(len(tracking[2]["evidence"]), 3)
            excel = pd.read_excel(io.BytesIO(saved["excel_bytes"]))
            self.assertEqual(excel["№"].tolist(), [2, 1, 3])
            self.assertEqual(excel["Статус аудита"].tolist(),
                             ["Проверка не выполнена", "ОК", "ОК"])
            self.assertEqual(excel["Где проверить"].tolist(),
                             ["Место не установлено; ссылка ВОР: стр. 2", "стр. 1, 2", "стр. 1, 2"])
            self.assertFalse(any(str(col).startswith("[ПД]") for col in excel.columns))
            word = docx.Document(io.BytesIO(saved["word_bytes"]))
            self.assertEqual([r.cells[0].text for r in word.tables[0].rows[1:]], ["2", "1", "3"])
            self.assertIn("стр. 1, 2", word.tables[0].rows[2].cells[3].text)
            self.assertNotIn("Simulated provider failure", str(excel.values))
            self.assertNotIn("Simulated provider failure", word.tables[0].rows[2].cells[3].text)
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()

    def test_missing_explicit_pdf_never_binds_sole_file(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "ghost.pdf", 1]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(table["Статус аудита"].tolist(), ["Ошибка"])
            self.assertIn("ghost.pdf", table["Детали расхождений"].iloc[0])
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()

    def test_continuation_row_folds_refs_into_owner(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1],
                 ["", "", "", "", "", "project.pdf", 2]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(len(table), 1)
            self.assertEqual(table["Статус аудита"].tolist(), ["ОК"])
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()

    def test_distinct_documents_are_not_shown_ready_then_silently_partially_audited(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            with fitz.open() as pdf:
                pdf.new_page()
                pdf.save(folder / "second.pdf")
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf;second.pdf", 1]])
            try:
                table = at.session_state["audit_results"]["vor.csv"]["df"]
                self.assertEqual(table["Статус аудита"].tolist(), ["Ошибка"])
                self.assertIn("несколько PDF", table["Детали расхождений"].iloc[0])
                preflight = next(frame.value for frame in at.dataframe
                                 if "Статус привязки" in frame.value.columns)
                self.assertTrue(preflight["Статус привязки"].iloc[0].startswith("🔴"))
            finally:
                for pdf in at.session_state["pdf_pool"].values():
                    pdf.close()

    def test_partly_invalid_page_list_cannot_be_ok(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            at = self._run_audit(
                Path(folder),
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", "1 999"]])
            try:
                table = at.session_state["audit_results"]["vor.csv"]["df"]
                self.assertEqual(table["Статус аудита"].tolist(), ["Ошибка"])
                self.assertIn("999", table["Детали расхождений"].iloc[0])
                self.assertEqual(table["[ПД] Найденные страницы"].tolist(), ["1, 2"])
            finally:
                for pdf in at.session_state["pdf_pool"].values():
                    pdf.close()

    def test_audit_reads_vor_page_link_column(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на чертежи",
                 "Наименование файла", "[ВОР] Ссылка на страницу"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(table["Статус аудита"].tolist(), ["ОК"])
            self.assertNotIn("Страница не указана", table["Детали расхождений"].iloc[0])
            self.assertEqual(table["[ВОР] Ссылка на страницу"].tolist(), [1])
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()

    def test_audit_rerun_overwrites_stale_results(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            folder = Path(folder)
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на чертежи",
                 "Наименование файла", "[ВОР] Ссылка на страницу",
                 "Статус аудита", "Детали расхождений",
                 "[ПД] Найденный объем (нейросеть)", "[ПД] Фактическая страница"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1,
                  "Технический сбой", "stale", "-", "-"]])
            saved = at.session_state["audit_results"]["vor.csv"]
            table = saved["df"]
            self.assertEqual(list(table.columns).count("Статус аудита"), 1)
            self.assertEqual(table["Статус аудита"].tolist(), ["ОК"])
            for pdf in at.session_state["pdf_pool"].values():
                pdf.close()


    def test_broken_preparation_does_not_abort_other_positions(self):
        with tempfile.TemporaryDirectory(prefix="smeta-audit-test-") as folder:
            at = self._run_audit(
                Path(folder),
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1],
                 [2, "Pump station", 1, "шт", "", "project.pdf", 2],
                 [3, "Copper cable", 45, "m", "", "project.pdf", 2]],
                preparation_fault_on_call=2)
            try:
                saved = at.session_state["audit_results"]["vor.csv"]
                table = saved["df"]
                self.assertEqual(table["Статус аудита"].tolist(),
                                 ["ОК", "Технический сбой", "ОК"])
                self.assertIn("Simulated preparation failure",
                              table["Детали расхождений"].iloc[1])
                self.assertEqual(table["[ПД] Запросов"].iloc[1], 0)
                self.assertEqual(table["[ПД] Поиск завершён"].iloc[1], "нет")
                self.assertEqual(saved["tracking"][1]["work_name"], "Pump station")
                self.assertIn("trace_id", saved["tracking"][1])
                excel = pd.read_excel(io.BytesIO(saved["excel_bytes"]))
                statuses = dict(zip(excel["№"], excel["Статус аудита"]))
                self.assertEqual(statuses, {1: "ОК", 2: "Проверка не выполнена",
                                            3: "ОК"})
            finally:
                for pdf in at.session_state["pdf_pool"].values():
                    pdf.close()

    def test_history_survives_ui_restart(self):
        import shutil

        from storage_blobs import FileBlobStore
        from storage_db import open_repository

        db_url = os.environ.get(
            "TEST_POSTGRES_URL",
            "postgresql+psycopg://nikgretenok@localhost:5432/smeta_test",
        )
        with tempfile.TemporaryDirectory(prefix="smeta-history-test-") as tmp:
            folder = Path(tmp) / "project"
            folder.mkdir()
            blobs = Path(tmp) / "blobs"
            blobs.mkdir()
            open_repository(db_url, blob_store=FileBlobStore(blobs)).reset_for_tests()
            at = self._run_audit(
                folder,
                ["№ п.п.", "Наименование работ, ресурсов, затрат по проекту",
                 "Объём работ / Количество", "Ед. изм.", "Ссылка на ПД",
                 "Наименование файла", "Номер страниц (через пробел)"],
                [[1, "Copper cable", 45, "m", "", "project.pdf", 1]],
                extra_env={"DATABASE_URL": db_url, "BLOB_STORE_URL": f"file://{blobs}"},
            )
            try:
                self.assertIn("run_id", at.session_state["audit_results"]["vor.csv"])
            finally:
                for pdf in at.session_state["pdf_pool"].values():
                    pdf.close()
            shutil.rmtree(folder)
            folder.mkdir()
            with patch.dict(
                os.environ,
                {
                    "YANDEX_CLOUD_API_KEY": "test-only",
                    "YANDEX_CLOUD_FOLDER": "test-folder",
                    "DATABASE_URL": db_url,
                    "BLOB_STORE_URL": f"file://{blobs}",
                },
            ), patch("audit_tracing._configured_client", return_value=None):
                restarted = AppTest.from_file(
                    str(Path(__file__).resolve().parents[1] / "app.py"),
                    default_timeout=30,
                )
                restarted.run()
                self.assertEqual(len(restarted.exception), 0, str(restarted.exception))
                folder_input = next(
                    item
                    for item in restarted.text_input
                    if item.label == "Путь к папке проекта на диске"
                )
                folder_input.set_value(str(folder)).run()
                self.assertEqual(len(restarted.exception), 0, str(restarted.exception))
                headers = [h.value for h in restarted.header]
                self.assertIn("🗂️ История запусков", headers)
                run_box = next(s for s in restarted.selectbox if s.label == "Запуск")
                self.assertEqual(len(run_box.options), 1)
                tables = [frame.value for frame in restarted.dataframe]
                self.assertTrue(
                    any(
                        "Статус аудита" in map(str, frame.columns)
                        and (frame["Статус аудита"] == "ОК").any()
                        for frame in tables
                        if hasattr(frame, "columns")
                    )
                )


if __name__ == "__main__":
    unittest.main()
