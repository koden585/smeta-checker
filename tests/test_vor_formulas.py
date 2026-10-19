"""VOR workbook quantity formulas: parsing, evaluation, provenance."""

import copy
import io
import sys
import unittest
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd
from audit_models import VorPosition
from vor_formula import (
    MAX_CELL_VISITS,
    MAX_FORMULA_LENGTH,
    MAX_REFERENCE_DEPTH,
    calculate_cell,
    normalize_address,
    snapshot_cells,
)
from vor_parse import load_vor_positions


def _cells(formula=None, cached=None, plain=None):
    entry = {}
    if formula is not None:
        entry["formula"] = formula
    entry["cached"] = cached
    entry["plain"] = plain
    return entry


class NormalizeAddressTests(unittest.TestCase):
    def test_pinned_and_case_insensitive(self):
        self.assertEqual(normalize_address("$d$23"), "D23")
        self.assertEqual(normalize_address("m23"), "M23")

    def test_rejects_non_addresses(self):
        for bad in ("", "D", "D0", "D1048577", "D23!", "A1:B2", "1D", None, 42):
            self.assertIsNone(normalize_address(bad))


class CalculateCellTests(unittest.TestCase):
    def test_workbook_case_number_six(self):
        cells = {
            "D18": _cells(plain=526, cached=526),
            "D22": _cells(formula="=ROUND(D18/(2.78*1.26)*1.1,0)", cached=165),
            "M23": _cells(formula="=ROUND(L18/L23/(L22/1000),0)", cached=5),
            "L18": _cells(plain=526, cached=526),
            "L22": _cells(plain=2780, cached=2780),
            "L23": _cells(plain=40, cached=40),
            "D23": _cells(
                formula="=ROUND(2*(M23*(D22/M23-1)+(M23-1)),0)", cached=328
            ),
        }
        result = calculate_cell("D23", cells)
        self.assertEqual(result["status"], "calculated")
        self.assertEqual(result["value"], "328")
        self.assertEqual(result["expression"], "=ROUND(2*(M23*(D22/M23-1)+(M23-1)),0)")
        self.assertEqual(result["cached_value"], 328)
        self.assertFalse(result["cached_mismatch"])
        self.assertIsNone(result["error"])
        deps = {d["cell"]: d for d in result["dependencies"] if d["origin"] == "vor_cell"}
        self.assertEqual(deps["D22"]["value"], "165")
        self.assertEqual(deps["D22"]["expression"], "=ROUND(D18/(2.78*1.26)*1.1,0)")
        self.assertEqual(deps["M23"]["value"], "5")

    def test_workbook_case_number_fifteen(self):
        cells = {
            "D30": _cells(plain=1879, cached=1879),
            "D33": _cells(formula="=D30*0.15", cached=281.84999999999997),
        }
        result = calculate_cell("D33", cells)
        self.assertEqual(result["status"], "calculated")
        self.assertEqual(Decimal(result["value"]), Decimal("281.85"))
        self.assertFalse(result["cached_mismatch"])

    def test_stale_cached_value_is_error_not_silent(self):
        cells = {"A1": _cells(formula="=2*3", cached=7)}
        result = calculate_cell("A1", cells)
        self.assertEqual(result["status"], "error")
        self.assertEqual(result["value"], "6")
        self.assertEqual(result["cached_value"], 7)
        self.assertTrue(result["cached_mismatch"])
        

    def test_missing_cached_value_computes_cleanly(self):
        cells = {"A1": _cells(formula="=2*3", cached=None, plain=None)}
        result = calculate_cell("A1", cells)
        self.assertEqual(result["status"], "calculated")
        self.assertEqual(result["value"], "6")
        self.assertFalse(result["cached_mismatch"])

    def test_cycle_is_explicit(self):
        cells = {
            "A1": _cells(formula="=B1+1", cached=2),
            "B1": _cells(formula="=A1+1", cached=3),
        }
        result = calculate_cell("A1", cells)
        self.assertEqual(result["status"], "error")
        

    def test_self_reference_is_a_cycle(self):
        result = calculate_cell("A1", {"A1": _cells(formula="=A1+1", cached=2)})
        self.assertEqual(result["status"], "error")
        

    def test_invalid_reference_is_error(self):
        result = calculate_cell("A1", {"A1": _cells(formula="=ZZZ9999999+1", cached=1)})
        self.assertEqual(result["status"], "error")
        

    def test_external_reference_is_unsupported(self):
        result = calculate_cell("A1", {"A1": _cells(formula="=Sheet2!A1+1", cached=1)})
        self.assertEqual(result["status"], "unsupported")
        

    def test_bracket_external_reference_is_unsupported(self):
        result = calculate_cell("A1", {"A1": _cells(formula="=[book.xlsx]S!A1", cached=1)})
        self.assertEqual(result["status"], "unsupported")

    def test_unsupported_function_is_unsupported(self):
        result = calculate_cell("A1", {"A1": _cells(formula="=SUM(A2:A3)", cached=1)})
        self.assertEqual(result["status"], "unsupported")
        

    def test_power_and_percent_are_unsupported(self):
        for expr in ("=2^3", "=50%"):
            result = calculate_cell(
                "A1", {"A1": _cells(formula=expr, cached=None, plain=None)}
            )
            self.assertEqual(result["status"], "unsupported", expr)

    def test_text_and_range_are_unsupported(self):
        for expr in ('="abc"+1', "=A2:A3"):
            result = calculate_cell(
                "A1", {"A1": _cells(formula=expr, cached=None, plain=None)}
            )
            self.assertEqual(result["status"], "unsupported", expr)

    def test_division_by_zero_is_error(self):
        result = calculate_cell("A1", {"A1": _cells(formula="=1/0", cached=0)})
        self.assertEqual(result["status"], "error")
        

    def test_non_numeric_reference_is_error(self):
        cells = {
            "A1": _cells(formula="=B1*2", cached=2),
            "B1": _cells(plain="текст", cached="текст"),
        }
        result = calculate_cell("A1", cells)
        self.assertEqual(result["status"], "error")
        

    def test_empty_reference_is_error(self):
        cells = {
            "A1": _cells(formula="=B1*2", cached=2),
            "B1": _cells(plain=None, cached=None),
        }
        result = calculate_cell("A1", cells)
        self.assertEqual(result["status"], "error")
        

    def test_negative_ties_round_half_up(self):
        cells = {"A1": _cells(formula="=ROUND(-2.5,0)", cached=-3)}
        result = calculate_cell("A1", cells)
        self.assertEqual(result["status"], "calculated")
        self.assertEqual(result["value"], "-3")

    def test_round_fractional_digits_rejected(self):
        result = calculate_cell(
            "A1", {"A1": _cells(formula="=ROUND(1.5,1.5)", cached=1.5)}
        )
        self.assertEqual(result["status"], "error")
        

    def test_deep_chain_is_bounded(self):
        cells = {}
        for i in range(1, MAX_REFERENCE_DEPTH + 5):
            addr = f"A{i}"
            nxt = f"A{i + 1}"
            cells[addr] = (
                _cells(formula=f"={nxt}+0", cached=1)
                if i < MAX_REFERENCE_DEPTH + 4
                else _cells(plain=1, cached=1)
            )
        result = calculate_cell("A1", cells)
        self.assertEqual(result["status"], "error")
        

    def test_too_many_references_are_bounded(self):
        refs = [f"B{i}" for i in range(1, MAX_CELL_VISITS + 2)]
        cells = {"A1": _cells(formula="=" + "+".join(refs), cached=1)}
        for ref in refs:
            cells[ref] = _cells(plain=1, cached=1)
        self.assertLessEqual(len(cells["A1"]["formula"]), MAX_FORMULA_LENGTH + 1)
        result = calculate_cell("A1", cells)
        self.assertEqual(result["status"], "error")
        

    def test_overlong_formula_is_error(self):
        result = calculate_cell(
            "A1", {"A1": _cells(formula="=" + "1+" * 300 + "1", cached=None)}
        )
        self.assertEqual(result["status"], "error")
        

    def test_huge_magnitude_is_error(self):
        result = calculate_cell(
            "A1", {"A1": _cells(formula="=9999999999999999*10", cached=None)}
        )
        self.assertEqual(result["status"], "error")
        

    def test_unary_signs_supported(self):
        result = calculate_cell(
            "A1", {"A1": _cells(formula="=---5+3", cached=-2)}
        )
        self.assertEqual(result["status"], "calculated")
        self.assertEqual(result["value"], "-2")

    def test_fractional_cell_reference_preserves_decimal_value(self):
        result = calculate_cell("A1", {"A1": _cells(formula="=1879*B1"),
                                       "B1": _cells(plain=0.15)})
        self.assertEqual(result["status"], "calculated")
        self.assertEqual(result["value"], "281.85")

    def test_deep_parentheses_return_explicit_error(self):
        result = calculate_cell("A1", {"A1": _cells(formula="=" + "(" * 100 + "1" + ")" * 100)})
        self.assertEqual(result["status"], "error")
        

    def test_stale_dependency_is_not_hidden_by_matching_final_cache(self):
        result = calculate_cell("A1", {"A1": _cells(formula="=0*B1+1", cached=1),
                                       "B1": _cells(formula="=1+1", cached=3)})
        self.assertEqual(result["status"], "error")
        

    def test_dependencies_traceable(self):
        cells = {
            "A1": _cells(formula="=B1*2+7", cached=13),
            "B1": _cells(plain=3, cached=3),
        }
        result = calculate_cell("A1", cells)
        self.assertEqual(result["status"], "calculated")
        self.assertEqual(result["value"], "13")
        by_origin = {}
        for dep in result["dependencies"]:
            by_origin.setdefault(dep["origin"], []).append(dep)
        self.assertEqual(by_origin["vor_cell"][0]["cell"], "B1")
        self.assertEqual(by_origin["vor_cell"][0]["expression"], "3")
        literals = {d["expression"] for d in by_origin["formula_literal"]}
        self.assertIn("2", literals)
        self.assertIn("7", literals)
        self.assertNotIn("formula", by_origin["vor_cell"][0])
    def test_snapshot_records_formulas_and_cached_values_independently(self):
        class Cell:
            def __init__(self, address, value):
                self.coordinate = address
                self.value = value

        class Sheet:
            def __init__(self, data):
                self._data = data

            def iter_rows(self):
                return [[Cell("A1", "=1+1")]]

            def __getitem__(self, address):
                return Cell(address, self._data[address])

        snapshot = snapshot_cells(Sheet({"A1": 2}), Sheet({"A1": 2}))
        self.assertEqual(snapshot["A1"]["formula"], "=1+1")
        self.assertEqual(snapshot["A1"]["cached"], 2)
        self.assertIsNone(snapshot["A1"]["plain"])





def _workbook_bytes(cells):
    import openpyxl

    book = openpyxl.Workbook()
    ws = book.active
    ws["A1"] = "№ п.п."
    ws["B1"] = "Наименование работ"
    ws["C1"] = "Тип позиции"
    ws["D1"] = "Объем работ / Количество"
    ws["E1"] = "Ед. изм."
    ws["F1"] = "Формула текст"
    for addr, value in cells.items():
        ws[addr] = value
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class VorParseCalculationTests(unittest.TestCase):
    HEADER = {
        "A1": "№ п.п.",
        "B1": "Наименование работ",
        "C1": "Тип позиции",
        "D1": "Объем работ / Количество",
        "E1": "Ед. изм.",
        "F1": "Формула текст",
    }

    def test_merged_header_and_fractional_dependency_are_preserved(self):
        import openpyxl
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.merge_cells("A1:F1")
        sheet["A1"] = "Расчётная ведомость"
        for cell, value in self.HEADER.items():
            sheet[cell.replace("1", "3")] = value
        for cell, value in {"A4": 1, "B4": "Укладка плит", "C4": "Работа", "D4": "=1879*D5",
                            "E4": "м3", "D5": 0.15}.items():
            sheet[cell] = value
        stream = io.BytesIO()
        book.save(stream)
        positions, _ = load_vor_positions(stream, "vor.xlsx")
        self.assertEqual(positions[0].volume, "281.85")
        self.assertEqual(positions[0].calculation["cell"], "D4")
        self.assertEqual(positions[0].calculation["status"], "calculated")

    def _positions(self, cells):
        merged = dict(self.HEADER)
        merged.update(cells)
        data = _workbook_bytes(merged)
        positions, _meta = load_vor_positions(io.BytesIO(data), "vor.xlsx")
        return positions

    def test_numeric_formula_evaluates_with_row_offset(self):
        positions = self._positions(
            {
                "A2": "Раздел",
                "A3": 1,
                "B3": "Скрепление плит",
                "C3": "Работа",
                "D3": "=ROUND(2*(M3*(D2/M3-1)+(M3-1)),0)",
                "F3": "2*(5*(165/5-1)+(5-1))",
                "D2": 165,
                "M3": 5,
            }
        )
        self.assertEqual(len(positions), 1)
        position = positions[0]
        self.assertEqual(position.number, "1")
        self.assertEqual(position.position_type, "Работа")
        self.assertEqual(position.formula, "2*(5*(165/5-1)+(5-1))")
        calc = position.calculation
        self.assertEqual(calc["status"], "calculated")
        self.assertEqual(calc["cell"], "D3")
        self.assertEqual(calc["value"], "328")
        self.assertIsNone(calc["error"])
        self.assertFalse(calc["cached_mismatch"])
        deps = {d["cell"]: d for d in calc["dependencies"] if d["cell"]}
        self.assertEqual(deps["D2"]["value"], "165")

    def test_uncached_formula_fills_volume(self):
        import openpyxl

        cells = dict(self.HEADER)
        cells.update(
            {
                "A2": 1,
                "B2": "Плиты",
                "C2": "Материал",
                "D2": "=D3*0.15",
                "E2": "м3",
                "D3": 1879,
            }
        )
        book = openpyxl.Workbook()
        ws = book.active
        for addr, value in cells.items():
            ws[addr] = value
        stream = io.BytesIO()
        book.save(stream)
        raw = stream.getvalue()
        cached_view = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        self.assertIsNone(cached_view.active["D2"].value)
        positions, _meta = load_vor_positions(io.BytesIO(raw), "vor.xlsx")
        self.assertEqual(positions[0].volume, "281.85")
        self.assertEqual(positions[0].calculation["status"], "calculated")

    def test_stale_cache_keeps_claimed_volume_with_mismatch(self):
        from vor_parse import _vor_positions

        cols = ["№ п.п.", "Наименование работ", "Ед. изм.",
                "Объем работ / Количество", "Тип позиции"]
        df = pd.DataFrame([["1", "Плиты", "шт", "7", "Материал"]], columns=cols)
        df.attrs["vor_header_idx"] = 0
        df.attrs["vor_cells"] = {
            "D2": {"formula": "=2*3", "cached": 7, "plain": None}
        }
        positions = _vor_positions(df)
        calc = positions[0].calculation
        self.assertEqual(calc["status"], "error")
        self.assertEqual(calc["value"], "6")
        self.assertTrue(calc["cached_mismatch"])
        # Claimed quantity stays the caller's; only the mismatch is exposed.
        self.assertEqual(positions[0].volume, "7")

    def test_filtered_row_offsets_map_to_original_excel_rows(self):
        from vor_parse import _vor_positions

        cols = ["№ п.п.", "Наименование работ", "Ед. изм.",
                "Объем работ / Количество", "Тип позиции"]
        df = pd.DataFrame(
            [["1", "A", "шт", "10", "Работа"], ["2", "B", "шт", "20", "Работа"]],
            columns=cols,
            index=[3, 7],
        )
        df.attrs["vor_header_idx"] = 14
        df.attrs["vor_cells"] = {
            "D19": {"formula": "=5*2", "cached": 10, "plain": None},
            "D23": {"formula": "=5*4", "cached": 20, "plain": None},
        }
        positions = _vor_positions(df)
        self.assertEqual(len(positions), 2)
        self.assertEqual(positions[0].calculation["cell"], "D19")
        self.assertEqual(positions[0].calculation["value"], "10")
        self.assertEqual(positions[1].calculation["cell"], "D23")
        self.assertEqual(positions[1].calculation["value"], "20")

    def test_csv_positions_keep_empty_calculation(self):
        csv = (
            "№ п.п.;Наименование работ;Объём работ / Количество;Ед. изм.;"
            "Тип позиции;Формула текст\n"
            "1;Copper cable;45;m;Работа;2*3\n"
        )

        class FakeFile(io.BytesIO):
            def getvalue(self):
                return super().getvalue()

        positions, _ = load_vor_positions(FakeFile(csv.encode("utf-8-sig")), "vor.csv")
        self.assertEqual(positions[0].calculation, {})
        self.assertEqual(positions[0].position_type, "Работа")
        self.assertEqual(positions[0].volume, "45")

    def test_position_type_defaults_empty(self):
        csv = "№ п.п.;Наименование работ;Объём работ / Количество;Ед. изм.\n1;Cable;45;m\n"

        class FakeFile(io.BytesIO):
            def getvalue(self):
                return super().getvalue()

        positions, _ = load_vor_positions(FakeFile(csv.encode("utf-8-sig")), "vor.csv")
        self.assertEqual(positions[0].position_type, "")

    def test_legacy_task_round_trip_preserves_new_fields(self):
        calc = calculate_cell(
            "A1", {"A1": _cells(formula="=2*3", cached=6, plain=None)}
        )
        position = VorPosition(
            index=0,
            work_name="Cable",
            volume="6",
            unit="m",
            position_type="Работа",
            calculation=calc,
        )
        task = position.legacy_task()
        self.assertEqual(task["position_type"], "Работа")
        self.assertEqual(task["calculation"]["value"], "6")
        task["calculation"]["value"] = "tampered"
        self.assertEqual(position.calculation["value"], "6")
        restored = VorPosition.from_legacy_task(copy.deepcopy(task))
        self.assertEqual(restored.position_type, "Работа")
        self.assertEqual(restored.calculation["value"], "tampered")

    def test_calculation_is_json_safe(self):
        import json

        positions = self._positions(
            {"A2": 1, "B2": "Cable", "C2": "Работа", "D2": "=2*3", "E2": "m"}
        )
        json.dumps(positions[0].calculation, ensure_ascii=False)


if __name__ == "__main__":
    unittest.main()
